"""
DXF RELAÇÃO DO AÇO → Excel
Extrai entidades MTEXT, classifica por X/Y, gera planilha formatada.
"""
import re
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from collections import defaultdict

DXF_FILE = r"C:\Users\PC2\Desktop\Drawing1.dxf"
OUT_FILE  = r"C:\Users\PC2\Desktop\Relacao_do_Aco.xlsx"

# ── column X ranges (derived from vertical LINE coordinates in the DXF) ──────
COL_X = {
    "AÇO":     (20.00, 20.90),
    "N":       (21.00, 21.70),
    "DIAM":    (21.70, 22.50),
    "QUANT":   (22.50, 23.30),
    "C.UNIT":  (23.30, 24.30),
    "C.TOTAL": (24.30, 25.50),
}
COLS = ["AÇO", "N", "DIAM", "QUANT", "C.UNIT", "C.TOTAL"]

# Y thresholds
Y_DATA_MAX    = 45.40   # rows below this are data
Y_SECTION_MIN = 45.85   # rows above this (and below title) are section labels
Y_TITLE_MIN   = 46.10   # rows above this are the title

def classify_col(x):
    for col, (lo, hi) in COL_X.items():
        if lo <= x < hi:
            return col
    return None

# ── parse MTEXT ───────────────────────────────────────────────────────────────
def parse_mtext(path):
    entities = []
    with open(path, encoding="ansi", errors="replace") as f:
        lines = f.readlines()
    i = 0
    while i < len(lines) - 1:
        if lines[i].strip() == "0" and lines[i+1].strip() == "MTEXT":
            entity = {}
            i += 2
            while i < len(lines) - 1:
                code = lines[i].strip()
                val  = lines[i+1].strip()
                if code == "10":
                    entity["x"] = float(val)
                elif code == "20":
                    entity["y"] = float(val)
                elif code == "1":
                    txt = re.sub(r'\{\\[^}]*\}', '', val)
                    txt = re.sub(r'\\[A-Za-z]+[0-9]*;?', '', txt)
                    entity["text"] = txt.strip()
                elif code == "0":
                    break
                i += 2
            if entity.get("text") and "x" in entity and "y" in entity:
                entities.append(entity)
        else:
            i += 1
    return entities

# ── group into table rows ─────────────────────────────────────────────────────
def build_table(entities):
    title_parts   = []
    section_labels = []   # list of (y, text)
    data_rows     = defaultdict(dict)  # y_key → {col: text}

    for e in entities:
        y = e["y"]
        x = e["x"]
        t = e["text"]

        if y >= Y_TITLE_MIN:
            title_parts.append(t)
        elif y >= Y_SECTION_MIN:
            section_labels.append((round(y, 2), t))
        elif y <= Y_DATA_MAX:
            col = classify_col(x)
            if col:
                data_rows[round(y, 2)][col] = t
            # else: ignore (lines outside column area)
        # else: column header region (45.40 < y < 45.85) — skip

    title = " — ".join(dict.fromkeys(title_parts)) if title_parts else "RELAÇÃO DO AÇO"
    # deduplicate & sort section labels descending
    seen = set()
    uniq_sections = []
    for y, t in sorted(section_labels, key=lambda x: -x[0]):
        if t not in seen:
            seen.add(t)
            uniq_sections.append((y, t))

    sorted_y = sorted(data_rows.keys(), reverse=True)
    return title, uniq_sections, sorted_y, data_rows

# ── Excel writer ──────────────────────────────────────────────────────────────
def write_excel(title, section_labels, sorted_y, data_rows, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relação do Aço"

    # ── styles ────────────────────────────────────────────────────────────────
    thin   = Side(style="thin")
    medium = Side(style="medium")
    bdr    = Border(left=thin, right=thin, top=thin, bottom=thin)
    bdr_m  = Border(left=medium, right=medium, top=medium, bottom=medium)
    center = Alignment(horizontal="center", vertical="center")
    left   = Alignment(horizontal="left", vertical="center")
    bold   = Font(bold=True, name="Calibri", size=10)
    normal = Font(name="Calibri", size=10)
    title_font   = Font(bold=True, name="Calibri", size=12)
    section_font = Font(bold=True, name="Calibri", size=10, italic=True)

    fill_header  = PatternFill("solid", fgColor="D9D9D9")
    fill_section = PatternFill("solid", fgColor="BDD7EE")
    fill_title   = PatternFill("solid", fgColor="1F3864")
    title_white  = Font(bold=True, name="Calibri", size=12, color="FFFFFF")
    fill_alt     = PatternFill("solid", fgColor="EBF3FB")

    # ── column widths ─────────────────────────────────────────────────────────
    col_widths = [10, 7, 8, 8, 10, 11]
    for i, w in enumerate(col_widths):
        ws.column_dimensions[get_column_letter(i+1)].width = w

    n = 6  # number of columns
    r = 1

    # ── title ─────────────────────────────────────────────────────────────────
    ws.merge_cells(f"A{r}:F{r}")
    c = ws[f"A{r}"]
    c.value = title
    c.font  = title_white
    c.alignment = center
    c.fill  = fill_title
    for col in range(1, n+1):
        ws.cell(row=r, column=col).border = bdr_m
    ws.row_dimensions[r].height = 20
    r += 1

    # ── section labels before the column headers ──────────────────────────────
    for _y, label in section_labels:
        ws.merge_cells(f"A{r}:F{r}")
        c = ws[f"A{r}"]
        c.value = label
        c.font  = section_font
        c.alignment = center
        c.fill  = fill_section
        for col in range(1, n+1):
            ws.cell(row=r, column=col).border = bdr
        r += 1

    # ── column headers (merged into 2 rows) ───────────────────────────────────
    header_vals = COLS
    sub_vals    = ["", "", "(mm)", "", "(cm)", "(cm)"]
    for i, (h, s) in enumerate(zip(header_vals, sub_vals)):
        if s:
            # Two sub-rows
            ws.cell(row=r,   column=i+1, value=h).font = bold
            ws.cell(row=r,   column=i+1).alignment = center
            ws.cell(row=r,   column=i+1).fill = fill_header
            ws.cell(row=r,   column=i+1).border = bdr
            ws.cell(row=r+1, column=i+1, value=s).font = Font(italic=True, name="Calibri", size=9)
            ws.cell(row=r+1, column=i+1).alignment = center
            ws.cell(row=r+1, column=i+1).fill = fill_header
            ws.cell(row=r+1, column=i+1).border = bdr
        else:
            # Merge two rows for this column
            ws.merge_cells(start_row=r, start_column=i+1, end_row=r+1, end_column=i+1)
            c = ws.cell(row=r, column=i+1, value=h)
            c.font = bold
            c.alignment = center
            c.fill = fill_header
            for row_off in [0, 1]:
                ws.cell(row=r+row_off, column=i+1).border = bdr
    r += 2

    # ── data rows ─────────────────────────────────────────────────────────────
    prev_aco = ""
    for idx, y_key in enumerate(sorted_y):
        row_data = data_rows[y_key]
        aco_val = row_data.get("AÇO", "")
        if aco_val:
            prev_aco = aco_val
        else:
            aco_val = prev_aco

        values = [
            aco_val,
            row_data.get("N", ""),
            row_data.get("DIAM", ""),
            row_data.get("QUANT", ""),
            row_data.get("C.UNIT", ""),
            row_data.get("C.TOTAL", ""),
        ]

        row_fill = fill_alt if idx % 2 == 1 else None

        for i, val in enumerate(values):
            cell = ws.cell(row=r, column=i+1)
            cell.border = bdr
            cell.alignment = center
            if row_fill:
                cell.fill = row_fill
            cell.font = normal
            if isinstance(val, str) and val not in ("VAR", ""):
                try:
                    num = float(val)
                    cell.value = int(num) if num == int(num) else num
                except ValueError:
                    cell.value = val
            else:
                cell.value = val

        r += 1

    wb.save(out_path)
    print(f"Salvo: {out_path}")
    print(f"Linhas de dados: {len(sorted_y)}")

# ── main ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("Lendo DXF...")
    entities = parse_mtext(DXF_FILE)
    print(f"  MTEXT encontrados: {len(entities)}")
    title, sections, sorted_y, data_rows = build_table(entities)
    print(f"  Título: {title!r}")
    print(f"  Seções: {[t for _,t in sections]}")
    print(f"  Linhas de dados: {len(sorted_y)}")
    write_excel(title, sections, sorted_y, data_rows, OUT_FILE)
