# -*- coding: utf-8 -*-
"""
DXF -> Excel  |  Relação do Aço  |  Marcos Roldão
=================================================
Converte tabelas de armação de DXF para planilha Excel formatada.

Ao executar sem argumentos, abre uma janela para escolher o arquivo DXF.
"""

import re
import sys
import unicodedata
import tkinter as tk

from pathlib import Path
from os import PathLike
from collections import defaultdict
from tkinter import filedialog, messagebox

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERRO: instale openpyxl → pip install openpyxl")
    sys.exit(1)


# ─────────────────────────────────────────────────────────────────────────────
#  PARSING DXF
# ─────────────────────────────────────────────────────────────────────────────

def _clean_mtext(text: str) -> str:
    """Remove códigos básicos de formatação MTEXT."""
    text = re.sub(r'\{\\[^}]*\}', '', text)
    text = re.sub(r'\\[A-Za-z]+[0-9]*;?', '', text)
    return text.strip()


def parse_text_entities(path: str) -> list[dict]:
    """
    Extrai entidades TEXT e MTEXT do DXF.
    Retorna lista de {"x": float, "y": float, "text": str}.
    """
    entities = []

    try:
        with open(path, encoding="utf-8-sig") as f:
            lines = f.readlines()
    except UnicodeDecodeError:
        with open(path, encoding="cp1252", errors="replace") as f:
            lines = f.readlines()

    i = 0

    while i < len(lines) - 1:
        entity_type = lines[i + 1].strip() if lines[i].strip() == "0" else None

        if entity_type in ("TEXT", "MTEXT"):
            e = {}
            text_parts = []
            i += 2

            while i < len(lines) - 1:
                code = lines[i].strip()
                val = lines[i + 1].strip()

                if code == "0":
                    break

                try:
                    if code == "10":
                        e["x"] = float(val)
                    elif code == "20":
                        e["y"] = float(val)
                except ValueError:
                    pass

                if code in ("1", "3"):
                    raw = _clean_mtext(val) if entity_type == "MTEXT" else val
                    if raw:
                        text_parts.append(raw)

                i += 2

            if text_parts:
                e["text"] = "".join(text_parts).strip()

            if e.get("text") and "x" in e and "y" in e:
                entities.append(e)

        else:
            i += 1

    return entities


# ─────────────────────────────────────────────────────────────────────────────
#  AUTO-DETECT TABLE STRUCTURE
# ─────────────────────────────────────────────────────────────────────────────

HEADER_KEYWORDS = {
    "n",
    "aco",
    "diam",
    "qtd",
    "quant",
    "c.unit",
    "c.total",
    "c.unit(cm)",
    "c.total(cm)",
}


def _ascii_norm(text: str) -> str:
    """Remove acentos e normaliza para comparação."""
    return (
        unicodedata.normalize("NFD", text)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
        .strip()
    )


def _is_header_keyword(text: str) -> bool:
    return _ascii_norm(text) in HEADER_KEYWORDS


def group_by_y(entities: list[dict], tolerance: float = 0.05) -> dict:
    """
    Agrupa entidades por coordenada Y.
    Retorna {y_key: {x_key: text}}.
    """
    rows = defaultdict(dict)

    for e in entities:
        y_key = round(e["y"] / tolerance) * tolerance
        x_key = round(e["x"], 3)
        rows[y_key][x_key] = e["text"]

    return rows


def find_header_row(rows: dict) -> tuple[float, dict] | tuple[None, None]:
    """
    Localiza a linha de cabeçalho.
    """
    sorted_ys = sorted(rows.keys(), reverse=True)

    for y_key in sorted_ys:
        row = rows[y_key]
        matches = sum(1 for t in row.values() if _is_header_keyword(t))

        if matches >= 2:
            merged = dict(row)

            for y2 in sorted_ys:
                if y2 == y_key:
                    continue

                if abs(y2 - y_key) <= 0.5:
                    row2 = rows[y2]
                    has_kw = any(_is_header_keyword(t) for t in row2.values())
                    is_units = all(
                        t in ("(mm)", "(cm)", "mm", "cm")
                        for t in row2.values()
                    )

                    if has_kw and not is_units:
                        merged.update(row2)

            return y_key, merged

    return None, None


def build_column_map(header_row: dict) -> tuple[list[str], list[tuple[float, float]]]:
    """
    Cria mapa de colunas baseado no X dos cabeçalhos.
    """
    sorted_items = sorted(header_row.items())
    col_names = [label for _, label in sorted_items]
    col_xs = [x for x, _ in sorted_items]

    boundaries = []

    for i, x in enumerate(col_xs):
        if i == 0:
            lo = x - 1.0
        else:
            lo = (col_xs[i - 1] + x) / 2.0

        if i == len(col_xs) - 1:
            hi = x + 2.0
        else:
            hi = (x + col_xs[i + 1]) / 2.0

        boundaries.append((lo, hi))

    return col_names, boundaries


def classify_x(
    x: float,
    col_names: list[str],
    x_ranges: list[tuple[float, float]],
) -> str | None:
    for col, (lo, hi) in zip(col_names, x_ranges):
        if lo <= x < hi:
            return col

    return None


# ─────────────────────────────────────────────────────────────────────────────
#  BUILD TABLE
# ─────────────────────────────────────────────────────────────────────────────

def build_table(entities: list[dict]) -> dict:
    rows = group_by_y(entities)
    header_y, header_row = find_header_row(rows)

    if header_y is None or header_row is None:
        raise ValueError("Não foi possível detectar cabeçalho de colunas no DXF.")

    col_names, x_ranges = build_column_map(header_row)

    title = ""
    sections = []
    data_rows = defaultdict(dict)

    all_above = [(y, rows[y]) for y in rows if y > header_y]
    max_y = max((y for y, _ in all_above), default=None)

    for y_key in sorted(rows.keys(), reverse=True):
        row = rows[y_key]

        if y_key > header_y:
            for _, text in sorted(row.items()):
                if _is_header_keyword(text):
                    continue

                if max_y and abs(y_key - max_y) < 0.01:
                    title = title or text
                else:
                    sections.append((y_key, text))

        elif abs(y_key - header_y) < 0.1:
            continue

        else:
            vals = list(row.values())

            if all(v in ("(mm)", "(cm)", "mm", "cm", "") for v in vals):
                continue

            for x_val, text in row.items():
                col = classify_x(x_val, col_names, x_ranges)

                if col:
                    data_rows[y_key][col] = text

    data_rows = {
        y: d
        for y, d in data_rows.items()
        if any(v not in ("", None) for v in d.values())
    }

    sorted_y = sorted(data_rows.keys(), reverse=True)

    seen_secs = set()
    uniq_sections = []

    for y, text in sorted(sections, key=lambda s: -s[0]):
        if text not in seen_secs and not _is_header_keyword(text):
            seen_secs.add(text)
            uniq_sections.append((y, text))

    return {
        "title": title or "RELAÇÃO DO AÇO",
        "sections": uniq_sections,
        "col_names": col_names,
        "sorted_y": sorted_y,
        "data_rows": data_rows,
    }


# ─────────────────────────────────────────────────────────────────────────────
#  EXCEL WRITER
# ─────────────────────────────────────────────────────────────────────────────

def to_number(val: str):
    """Converte texto para número quando possível."""
    if not isinstance(val, str):
        return val

    if val.strip() in ("VAR", ""):
        return val

    try:
        f = float(val.replace(",", "."))
        return int(f) if f == int(f) else f
    except (ValueError, OverflowError):
        return val


def write_excel(table: dict, out_path: str) -> None:
    title = table["title"]
    sections = table["sections"]
    col_names = table["col_names"]
    sorted_y = table["sorted_y"]
    data_rows = table["data_rows"]
    n_cols = len(col_names)

    wb = openpyxl.Workbook()

    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Relação do Aço")

    ws.title = "Relação do Aço"

    thin = Side(style="thin")
    medium = Side(style="medium")

    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    bdr_t = Border(left=medium, right=medium, top=medium, bottom=medium)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c_title = Font(bold=True, name="Calibri", size=12, color="FFFFFF")
    c_section = Font(bold=True, name="Calibri", size=10, italic=True)
    c_header = Font(bold=True, name="Calibri", size=10)
    c_normal = Font(name="Calibri", size=10)

    f_title = PatternFill("solid", fgColor="1F3864")
    f_section = PatternFill("solid", fgColor="BDD7EE")
    f_header = PatternFill("solid", fgColor="D9D9D9")
    f_alt = PatternFill("solid", fgColor="EBF3FB")

    base_w = max(8, 60 // n_cols)

    for i in range(n_cols):
        ws.column_dimensions[get_column_letter(i + 1)].width = base_w

    r = 1

    def merge_row(row: int, text: str, font: Font, fill: PatternFill, height: int = 18):
        ws.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=n_cols,
        )

        cell = ws.cell(row=row, column=1, value=text)
        cell.font = font
        cell.alignment = center
        cell.fill = fill

        for col in range(1, n_cols + 1):
            ws.cell(row=row, column=col).border = bdr_t if font == c_title else bdr

        ws.row_dimensions[row].height = height

    merge_row(r, title, c_title, f_title, height=22)
    r += 1

    for _, label in sections:
        merge_row(r, label, c_section, f_section)
        r += 1

    for i, col in enumerate(col_names):
        cell = ws.cell(row=r, column=i + 1, value=col)
        cell.font = c_header
        cell.alignment = center
        cell.fill = f_header
        cell.border = bdr

    ws.row_dimensions[r].height = 16
    r += 1

    first_col = col_names[0]
    prev_first = ""

    for idx, y_key in enumerate(sorted_y):
        row_data = data_rows[y_key]
        fill = f_alt if idx % 2 == 1 else None

        val0 = row_data.get(first_col, "")

        if val0:
            prev_first = val0
        elif _ascii_norm(first_col) in ("aco",):
            val0 = prev_first

        for i, col in enumerate(col_names):
            raw = val0 if i == 0 else row_data.get(col, "")
            val = to_number(raw)

            cell = ws.cell(row=r, column=i + 1, value=val)
            cell.font = c_normal
            cell.alignment = center
            cell.border = bdr

            if fill:
                cell.fill = fill

        r += 1

    wb.save(out_path)
    print(f"  OK  {out_path}  ({len(sorted_y)} linhas, {n_cols} colunas)")


# ─────────────────────────────────────────────────────────────────────────────
#  INTERFACE DE SELEÇÃO
# ─────────────────────────────────────────────────────────────────────────────

def select_dxf_file() -> str | None:
    downloads = Path.home() / "Downloads"
    initial_dir = downloads if downloads.exists() else Path.home()

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    file_path = filedialog.askopenfilename(
        title="Selecione o arquivo DXF",
        initialdir=str(initial_dir),
        filetypes=[
            ("Arquivos DXF", "*.dxf"),
            ("Todos os arquivos", "*.*"),
        ],
    )

    root.destroy()

    if not file_path:
        return None

    return file_path


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────────────────────────────────────

def process(
    dxf_path: str | PathLike[str],
    xlsx_path: str | PathLike[str] | None = None,
) -> None:
    dxf_path = str(dxf_path)

    if xlsx_path is not None:
        xlsx_path = str(xlsx_path)
    else:
        xlsx_path = str(Path(dxf_path).with_suffix(".xlsx"))

    print(f"\n>> {dxf_path}")

    entities = parse_text_entities(dxf_path)
    print(f"  Entidades de texto: {len(entities)}")

    table = build_table(entities)

    print(f"  Título   : {table['title']!r}")
    print(f"  Colunas  : {table['col_names']}")
    print(f"  Seções   : {[t for _, t in table['sections']]}")
    print(f"  Dados    : {len(table['sorted_y'])} linhas")

    write_excel(table, xlsx_path)


def main() -> None:
    args = sys.argv[1:]

    if len(args) == 0:
        dxf_file = select_dxf_file()

        if not dxf_file:
            print("Nenhum arquivo selecionado.")
            sys.exit(0)

        try:
            process(dxf_file)
            messagebox.showinfo(
                "Concluído",
                "Arquivo Excel gerado com sucesso.",
            )
        except Exception as ex:
            print(f"ERRO: {ex}")
            messagebox.showerror(
                "Erro ao processar DXF",
                str(ex),
            )
            sys.exit(1)

    elif len(args) == 1:
        process(args[0])

    elif len(args) == 2:
        process(args[0], args[1])

    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()