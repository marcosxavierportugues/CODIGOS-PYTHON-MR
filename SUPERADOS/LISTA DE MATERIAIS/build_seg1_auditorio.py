import pandas as pd
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

file_in  = r'D:\00-Obras - SLN\IFMT - SLN\00- ENVIADO IFMT\SEG 1\SEG 1- LAJES AUDITORIO- ENVIADOS AO IF.xlsx'
file_out = r'D:\00-Obras - SLN\IFMT - SLN\00- ENVIADO IFMT\SEG 1\SEG 1- LAJES AUDITORIO - ORGANIZADO-BITOLA.xlsx'
TITULO   = 'SEG 1 LAJES AUDITORIO'

df = pd.read_excel(file_in, sheet_name='Sheet1', header=0)
df.columns = [str(c) for c in df.columns]

COL_NAME  = 0
COL_LAYER = 6
item_pairs = [(58,59),(60,61),(62,63),(64,65),(66,67),(68,69),(70,71),(72,73),(74,75),(76,77),(78,79)]


def parse_qty(val):
    if pd.isna(val): return None
    m = re.search(r'([\d\.]+)', str(val))
    return float(m.group(1)) if m else None

def extract_bitola(desc):
    m = re.search(r'([\d\.]+)\s*mm', str(desc))
    return float(m.group(1)) if m else None

def classify(desc):
    d = str(desc)
    if 'Forma' in d:          return 'FORMA'
    if 'Concreto' in d:       return 'CONCRETO'
    if 'Chapa' in d:          return 'CHAPA'
    if 'pun' in d.lower():    return 'PUNCAO'
    if 'Armadura' in d:       return 'ACO'
    return None


# Collect bitolas
all_bitolas = set()
has_puncao = has_chapa = False
for _, row in df.iterrows():
    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        if pd.notna(desc):
            cat = classify(desc)
            if cat == 'ACO':
                b = extract_bitola(desc)
                if b: all_bitolas.add(b)
            elif cat == 'PUNCAO': has_puncao = True
            elif cat == 'CHAPA':  has_chapa  = True

bitolas = sorted(all_bitolas)
print('Bitolas CA50:', bitolas)
print(f'Puncao: {has_puncao} | Chapa: {has_chapa}')

# Build records
records = []
for _, row in df.iterrows():
    elem  = str(row.iloc[COL_NAME])
    layer = str(row.iloc[COL_LAYER])
    concreto = forma = puncao = chapa = 0.0
    aco = {b: 0.0 for b in bitolas}

    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        if pd.isna(desc): continue
        cat = classify(desc)
        qty = parse_qty(row.iloc[qc])
        if qty is None: continue
        if cat == 'CONCRETO':      concreto += qty
        elif cat == 'FORMA':       forma    += qty
        elif cat == 'ACO':
            b = extract_bitola(desc)
            if b and b in aco: aco[b] += qty
        elif cat == 'PUNCAO':      puncao   += qty
        elif cat == 'CHAPA':       chapa    += qty

    rec = {'LAJE': elem, 'CAMADA': layer, 'CONCRETO_m3': concreto, 'FORMA_m2': forma}
    for b in bitolas: rec[f'ACO_{b}mm_kg'] = aco[b]
    if has_puncao: rec['PUNCAO_CA25_un'] = puncao
    if has_chapa:  rec['CHAPA_RETANG_un'] = chapa
    records.append(rec)

result = pd.DataFrame(records)

print('\n=== VALIDACAO ===')
print(f'CONCRETO: {result["CONCRETO_m3"].sum():.4f} m3')
print(f'FORMA:    {result["FORMA_m2"].sum():.4f} m2')
for b in bitolas:
    t = result[f'ACO_{b}mm_kg'].sum()
    if t > 0: print(f'ACO O{b}mm: {t:.4f} kg')
if has_puncao: print(f'PUNCAO CA25 O8mm: {result["PUNCAO_CA25_un"].sum():.0f} un')
if has_chapa:  print(f'CHAPA RETANG:     {result["CHAPA_RETANG_un"].sum():.0f} un')
print(f'Total lajes: {len(result)}')

# ======================== EXCEL ========================
wb = openpyxl.Workbook()
thin = Side(style='thin', color='BBBBBB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center  = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_al = Alignment(horizontal='left',   vertical='center')
right_al= Alignment(horizontal='right',  vertical='center')
wf = Font(name='Arial', bold=True, color='FFFFFF', size=10)
nf = Font(name='Arial', size=10)
hdr_blue   = PatternFill('solid', start_color='1F4E79')
hdr_green  = PatternFill('solid', start_color='375623')
hdr_purple = PatternFill('solid', start_color='4A235A')
hdr_red    = PatternFill('solid', start_color='7B1818')
hdr_brown  = PatternFill('solid', start_color='5C3317')
fill_total = PatternFill('solid', start_color='1F4E79')
fill_gold  = PatternFill('solid', start_color='D4AC0D')
fill_conc  = PatternFill('solid', start_color='D6E4F0')
fill_forma = PatternFill('solid', start_color='D5E8D4')
fill_aco   = PatternFill('solid', start_color='EDE7F6')
fill_punc  = PatternFill('solid', start_color='FDEBD0')

# --- ABA 1 DETALHADO ---
ws = wb.active
ws.title = 'DETALHADO'
hdr_lbls  = ['LAJE', 'CAMADA', 'CONCRETO', 'FORMA'] + [f'ACO O{b}mm' for b in bitolas]
unit_lbls = ['', '', 'm3', 'm2'] + ['kg'] * len(bitolas)
fills_h   = [hdr_blue, hdr_blue, hdr_blue, hdr_green] + [hdr_purple] * len(bitolas)
if has_puncao: hdr_lbls.append('PUNCAO CA25 O8mm'); unit_lbls.append('un'); fills_h.append(hdr_red)
if has_chapa:  hdr_lbls.append('CHAPA RETANG 16.2x3.2'); unit_lbls.append('un'); fills_h.append(hdr_brown)

for ci,(h,u,f) in enumerate(zip(hdr_lbls, unit_lbls, fills_h), 1):
    for ri,val in [(1,h),(2,u)]:
        c = ws.cell(row=ri, column=ci, value=val)
        c.font=wf; c.fill=f; c.alignment=center; c.border=border

for ri, rec in enumerate(records, 3):
    rf = PatternFill('solid', start_color='F5F5F5') if ri%2==1 else PatternFill('solid', start_color='FFFFFF')
    vals = [rec['LAJE'], rec['CAMADA'], rec['CONCRETO_m3'], rec['FORMA_m2']] + \
           [rec[f'ACO_{b}mm_kg'] for b in bitolas]
    if has_puncao: vals.append(rec['PUNCAO_CA25_un'])
    if has_chapa:  vals.append(rec['CHAPA_RETANG_un'])
    for ci,val in enumerate(vals, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.font=nf; c.border=border; c.fill=rf
        c.alignment = left_al if ci<=2 else right_al
        if ci>2: c.number_format='#,##0.00'

ri_tot = 3+len(records)
tvs = ['TOTAL','',result['CONCRETO_m3'].sum(),result['FORMA_m2'].sum()] + \
      [result[f'ACO_{b}mm_kg'].sum() for b in bitolas]
if has_puncao: tvs.append(result['PUNCAO_CA25_un'].sum())
if has_chapa:  tvs.append(result['CHAPA_RETANG_un'].sum())
for ci,v in enumerate(tvs, 1):
    c = ws.cell(row=ri_tot, column=ci, value=v)
    c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10); c.fill=fill_total; c.border=border
    c.alignment = left_al if ci<=2 else right_al
    if ci>2 and isinstance(v,float): c.number_format='#,##0.00'

extra = (1 if has_puncao else 0)+(1 if has_chapa else 0)
col_widths = [14,20,14,14]+[13]*len(bitolas)+[18]*extra
for ci,w in enumerate(col_widths,1): ws.column_dimensions[get_column_letter(ci)].width=w
ws.row_dimensions[1].height=35; ws.row_dimensions[2].height=18
ws.freeze_panes='A3'

# --- ABA 2 RESUMO TOTAL ---
ws2 = wb.create_sheet('RESUMO TOTAL')
ws2['A1'] = f'RESUMO GERAL - {TITULO}'
ws2['A1'].font = Font(name='Arial',bold=True,size=14,color='1F4E79')
ws2.merge_cells('A1:D1'); ws2['A1'].alignment=center
for ci,h in enumerate(['CATEGORIA','DESCRICAO','TOTAL','UNIDADE'],1):
    c=ws2.cell(row=3,column=ci,value=h); c.font=wf; c.fill=hdr_blue; c.alignment=center; c.border=border

rows_r = [('CONCRETO','Concreto C-30 - Abatimento 5 cm',result['CONCRETO_m3'].sum(),'m3'),
          ('FORMA','Forma - Estrutura - Concreto',result['FORMA_m2'].sum(),'m2')]
for b in bitolas:
    t = result[f'ACO_{b}mm_kg'].sum()
    if t>0: rows_r.append(('ACO',f'Armadura CA50 - O {b} mm',t,'kg'))
if has_puncao: rows_r.append(('PUNCAO','Armadura de puncao - CA25 - O 8.0 mm (9 cm)',result['PUNCAO_CA25_un'].sum(),'un'))
if has_chapa:  rows_r.append(('CHAPA','Armadura de puncao - Chapa retangular 16.2x3.2',result['CHAPA_RETANG_un'].sum(),'un'))
fc={'CONCRETO':fill_conc,'FORMA':fill_forma,'ACO':fill_aco,'PUNCAO':fill_punc,'CHAPA':fill_punc}
for ri,(cat,desc,val,unit) in enumerate(rows_r,4):
    f=fc.get(cat,fill_conc)
    for ci,v in enumerate([cat,desc,val,unit],1):
        c=ws2.cell(row=ri,column=ci,value=v); c.font=nf; c.border=border; c.fill=f
        c.alignment=left_al if ci==2 else center
        if ci==3: c.number_format='#,##0.00'

r_gt=4+len(rows_r)
ag=sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
for ci,v in enumerate(['ACO TOTAL','Total Geral ACO CA50 (todas as bitolas)',ag,'kg'],1):
    c=ws2.cell(row=r_gt,column=ci,value=v)
    c.font=Font(name='Arial',bold=True,size=10); c.fill=fill_gold; c.border=border
    c.alignment=left_al if ci==2 else center
    if ci==3: c.number_format='#,##0.00'
ws2.column_dimensions['A'].width=14; ws2.column_dimensions['B'].width=50
ws2.column_dimensions['C'].width=16; ws2.column_dimensions['D'].width=12

# --- ABA 3 RESUMO POR CAMADA ---
ws3 = wb.create_sheet('RESUMO POR CAMADA')
nc3 = 3+len(bitolas)+1+extra
ws3['A1'] = f'RESUMO POR CAMADA - {TITULO}'
ws3['A1'].font=Font(name='Arial',bold=True,size=14,color='1F4E79')
ws3.merge_cells(f'A1:{get_column_letter(nc3)}1'); ws3['A1'].alignment=center
h3=['CAMADA','CONCRETO (m3)','FORMA (m2)']+[f'ACO O{b}mm (kg)' for b in bitolas]+['ACO TOTAL (kg)']
if has_puncao: h3.append('PUNCAO CA25 (un)')
if has_chapa:  h3.append('CHAPA RETANG (un)')
for ci,h in enumerate(h3,1):
    c=ws3.cell(row=3,column=ci,value=h); c.font=wf; c.fill=hdr_blue; c.alignment=center; c.border=border

for ri,cam in enumerate(result['CAMADA'].unique(),4):
    sub=result[result['CAMADA']==cam]
    at=sum(sub[f'ACO_{b}mm_kg'].sum() for b in bitolas)
    v3=[cam,sub['CONCRETO_m3'].sum(),sub['FORMA_m2'].sum()]+\
       [sub[f'ACO_{b}mm_kg'].sum() for b in bitolas]+[at]
    if has_puncao: v3.append(sub['PUNCAO_CA25_un'].sum())
    if has_chapa:  v3.append(sub['CHAPA_RETANG_un'].sum())
    rf=PatternFill('solid',start_color='EBF3FB') if ri%2==0 else PatternFill('solid',start_color='FFFFFF')
    for ci,v in enumerate(v3,1):
        c=ws3.cell(row=ri,column=ci,value=v); c.font=nf; c.border=border; c.fill=rf
        c.alignment=left_al if ci==1 else right_al
        if ci>1: c.number_format='#,##0.00'

rtt=4+len(result['CAMADA'].unique())
aa=sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
tt=['TOTAL GERAL',result['CONCRETO_m3'].sum(),result['FORMA_m2'].sum()]+\
   [result[f'ACO_{b}mm_kg'].sum() for b in bitolas]+[aa]
if has_puncao: tt.append(result['PUNCAO_CA25_un'].sum())
if has_chapa:  tt.append(result['CHAPA_RETANG_un'].sum())
for ci,v in enumerate(tt,1):
    c=ws3.cell(row=rtt,column=ci,value=v)
    c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10); c.fill=fill_total; c.border=border
    c.alignment=left_al if ci==1 else right_al
    if ci>1: c.number_format='#,##0.00'
for ci,w in enumerate([22,16,16]+[13]*len(bitolas)+[16]+[18]*extra,1):
    ws3.column_dimensions[get_column_letter(ci)].width=w

wb.save(file_out)
print(f'\nArquivo salvo: {file_out}')
