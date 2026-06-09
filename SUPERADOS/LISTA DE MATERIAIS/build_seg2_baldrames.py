import pandas as pd
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

file_in  = r'D:\00-Obras - SLN\IFMT - SLN\00- ENVIADO IFMT\SEG 2\SEG 2 - VIGAS BALDRAMES  - ENVIADO AO IF.xlsx'
file_out = r'D:\00-Obras - SLN\IFMT - SLN\00- ENVIADO IFMT\SEG 2\SEG 2 - VIGAS BALDRAMES - ORGANIZADO-BITOLA.xlsx'

df = pd.read_excel(file_in, sheet_name='Sheet1', header=0)
df.columns = [str(c) for c in df.columns]

COL_NAME  = 0
COL_LAYER = 1
item_pairs = [(3,4),(5,6),(7,8),(9,10),(11,12),(13,14),(15,16),(17,18),(19,20),(21,22),(23,24)]


def parse_qty(val):
    if pd.isna(val):
        return None
    m = re.search(r'([\d\.]+)', str(val))
    return float(m.group(1)) if m else None


def extract_bitola(desc):
    m = re.search(r'([\d\.]+)\s*mm', str(desc))
    return float(m.group(1)) if m else None


def classify(desc):
    d = str(desc)
    if 'Forma' in d:
        return 'FORMA'
    if 'Concreto' in d:
        return 'CONCRETO'
    if 'Armadura' in d or 'rmadura' in d:
        return 'ACO'
    return None


# Collect all bitolas
all_bitolas = set()
for _, row in df.iterrows():
    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        if pd.notna(desc) and classify(desc) == 'ACO':
            b = extract_bitola(desc)
            if b:
                all_bitolas.add(b)

bitolas = sorted(all_bitolas)
print('Bitolas encontradas:', bitolas)

# Build structured records
records = []
for _, row in df.iterrows():
    elem = str(row.iloc[COL_NAME])
    layer = str(row.iloc[COL_LAYER])

    concreto = 0.0
    forma = 0.0
    aco = {b: 0.0 for b in bitolas}

    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        qty_raw = row.iloc[qc]
        if pd.isna(desc):
            continue
        cat = classify(desc)
        qty = parse_qty(qty_raw)
        if qty is None:
            continue
        if cat == 'CONCRETO':
            concreto += qty
        elif cat == 'FORMA':
            forma += qty
        elif cat == 'ACO':
            b = extract_bitola(desc)
            if b and b in aco:
                aco[b] += qty

    rec = {'VIGA': elem, 'CAMADA': layer, 'CONCRETO_m3': concreto, 'FORMA_m2': forma}
    for b in bitolas:
        rec[f'ACO_{b}mm_kg'] = aco[b]
    records.append(rec)

result = pd.DataFrame(records)

# VALIDATION
print('\n=== VALIDACAO - TOTAIS GERAIS ===')
print(f'Total CONCRETO: {result["CONCRETO_m3"].sum():.4f} m3')
print(f'Total FORMA:    {result["FORMA_m2"].sum():.4f} m2')
for b in bitolas:
    t = result[f'ACO_{b}mm_kg'].sum()
    if t > 0:
        print(f'Total ACO O{b}mm: {t:.4f} kg')
print(f'Total vigas/baldrames: {len(result)}')

# ========================
# BUILD EXCEL
# ========================
wb = openpyxl.Workbook()

thin = Side(style='thin', color='BBBBBB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_al = Alignment(horizontal='left', vertical='center')
right_al = Alignment(horizontal='right', vertical='center')

white_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
normal_font = Font(name='Arial', size=10)

hdr_blue   = PatternFill('solid', start_color='1F4E79')
hdr_green  = PatternFill('solid', start_color='375623')
hdr_purple = PatternFill('solid', start_color='4A235A')
fill_total = PatternFill('solid', start_color='1F4E79')
fill_gold  = PatternFill('solid', start_color='D4AC0D')
fill_conc  = PatternFill('solid', start_color='D6E4F0')
fill_forma = PatternFill('solid', start_color='D5E8D4')
fill_aco   = PatternFill('solid', start_color='EDE7F6')

# =============================================
# ABA 1 - DETALHADO
# =============================================
ws = wb.active
ws.title = 'DETALHADO'

header_labels = ['VIGA/BALDRAME', 'CAMADA', 'CONCRETO', 'FORMA'] + [f'ACO O{b}mm' for b in bitolas]
unit_labels   = ['', '', 'm3', 'm2'] + ['kg'] * len(bitolas)
fills_hdr     = [hdr_blue, hdr_blue, hdr_blue, hdr_green] + [hdr_purple] * len(bitolas)

for ci, (h, u, f) in enumerate(zip(header_labels, unit_labels, fills_hdr), 1):
    for ri, val in [(1, h), (2, u)]:
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = white_font
        c.fill = f
        c.alignment = center
        c.border = border

for ri, rec in enumerate(records, 3):
    row_fill = PatternFill('solid', start_color='F5F5F5') if ri % 2 == 1 else PatternFill('solid', start_color='FFFFFF')
    vals = [rec['VIGA'], rec['CAMADA'], rec['CONCRETO_m3'], rec['FORMA_m2']] + [rec[f'ACO_{b}mm_kg'] for b in bitolas]
    for ci, val in enumerate(vals, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = normal_font
        c.border = border
        c.fill = row_fill
        c.alignment = left_al if ci <= 2 else right_al
        if ci > 2:
            c.number_format = '#,##0.00'

ri_tot = 3 + len(records)
tot_vals = ['TOTAL', ''] + \
    [result['CONCRETO_m3'].sum(), result['FORMA_m2'].sum()] + \
    [result[f'ACO_{b}mm_kg'].sum() for b in bitolas]
for ci, v in enumerate(tot_vals, 1):
    c = ws.cell(row=ri_tot, column=ci, value=v)
    c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    c.fill = fill_total
    c.border = border
    c.alignment = left_al if ci <= 2 else right_al
    if ci > 2 and isinstance(v, float):
        c.number_format = '#,##0.00'

col_widths = [16, 20, 14, 14] + [13] * len(bitolas)
for ci, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 18
ws.freeze_panes = 'A3'

# =============================================
# ABA 2 - RESUMO TOTAL
# =============================================
ws2 = wb.create_sheet('RESUMO TOTAL')
ws2['A1'] = 'RESUMO GERAL - SEG 2 VIGAS BALDRAMES'
ws2['A1'].font = Font(name='Arial', bold=True, size=14, color='1F4E79')
ws2.merge_cells('A1:D1')
ws2['A1'].alignment = center

for ci, h in enumerate(['CATEGORIA', 'DESCRICAO', 'TOTAL', 'UNIDADE'], 1):
    c = ws2.cell(row=3, column=ci, value=h)
    c.font = white_font
    c.fill = hdr_blue
    c.alignment = center
    c.border = border

rows_resumo = [('CONCRETO', 'Concreto C-30 - Abatimento 12 cm', result['CONCRETO_m3'].sum(), 'm3'),
               ('FORMA', 'Forma - Estrutura - Concreto', result['FORMA_m2'].sum(), 'm2')]
for b in bitolas:
    t = result[f'ACO_{b}mm_kg'].sum()
    if t > 0:
        steel = 'CA60' if b == 5.0 else 'CA50'
        rows_resumo.append(('ACO', f'Armadura {steel} - O {b} mm', t, 'kg'))

fills_cat = {'CONCRETO': fill_conc, 'FORMA': fill_forma, 'ACO': fill_aco}
for ri, (cat, desc, val, unit) in enumerate(rows_resumo, 4):
    f = fills_cat.get(cat, fill_conc)
    for ci, v in enumerate([cat, desc, val, unit], 1):
        c = ws2.cell(row=ri, column=ci, value=v)
        c.font = normal_font
        c.border = border
        c.fill = f
        c.alignment = left_al if ci == 2 else center
        if ci == 3:
            c.number_format = '#,##0.00'

r_gt = 4 + len(rows_resumo)
aco_grand = sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
for ci, v in enumerate(['ACO TOTAL', 'Total Geral de Aco (todas as bitolas)', aco_grand, 'kg'], 1):
    c = ws2.cell(row=r_gt, column=ci, value=v)
    c.font = Font(name='Arial', bold=True, size=10)
    c.fill = fill_gold
    c.border = border
    c.alignment = left_al if ci == 2 else center
    if ci == 3:
        c.number_format = '#,##0.00'

ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 44
ws2.column_dimensions['C'].width = 16
ws2.column_dimensions['D'].width = 12

# =============================================
# ABA 3 - RESUMO POR CAMADA
# =============================================
ws3 = wb.create_sheet('RESUMO POR CAMADA')
n_cols_3 = 3 + len(bitolas) + 1
ws3['A1'] = 'RESUMO POR CAMADA - SEG 2 VIGAS BALDRAMES'
ws3['A1'].font = Font(name='Arial', bold=True, size=14, color='1F4E79')
ws3.merge_cells(f'A1:{get_column_letter(n_cols_3)}1')
ws3['A1'].alignment = center

hdrs3 = ['CAMADA', 'CONCRETO (m3)', 'FORMA (m2)'] + [f'ACO O{b}mm (kg)' for b in bitolas] + ['ACO TOTAL (kg)']
for ci, h in enumerate(hdrs3, 1):
    c = ws3.cell(row=3, column=ci, value=h)
    c.font = white_font
    c.fill = hdr_blue
    c.alignment = center
    c.border = border

camadas = result['CAMADA'].unique()
for ri, cam in enumerate(camadas, 4):
    sub = result[result['CAMADA'] == cam]
    aco_tot = sum(sub[f'ACO_{b}mm_kg'].sum() for b in bitolas)
    vals3 = [cam, sub['CONCRETO_m3'].sum(), sub['FORMA_m2'].sum()] + \
            [sub[f'ACO_{b}mm_kg'].sum() for b in bitolas] + [aco_tot]
    row_fill = PatternFill('solid', start_color='EBF3FB') if ri % 2 == 0 else PatternFill('solid', start_color='FFFFFF')
    for ci, v in enumerate(vals3, 1):
        c = ws3.cell(row=ri, column=ci, value=v)
        c.font = normal_font
        c.border = border
        c.fill = row_fill
        c.alignment = left_al if ci == 1 else right_al
        if ci > 1:
            c.number_format = '#,##0.00'

ri_tot3 = 4 + len(camadas)
aco_all = sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
totals3 = ['TOTAL GERAL', result['CONCRETO_m3'].sum(), result['FORMA_m2'].sum()] + \
          [result[f'ACO_{b}mm_kg'].sum() for b in bitolas] + [aco_all]
for ci, v in enumerate(totals3, 1):
    c = ws3.cell(row=ri_tot3, column=ci, value=v)
    c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    c.fill = fill_total
    c.border = border
    c.alignment = left_al if ci == 1 else right_al
    if ci > 1:
        c.number_format = '#,##0.00'

for ci, w in enumerate([22, 16, 16] + [14] * len(bitolas) + [16], 1):
    ws3.column_dimensions[get_column_letter(ci)].width = w

wb.save(file_out)
print(f'\nArquivo salvo: {file_out}')
