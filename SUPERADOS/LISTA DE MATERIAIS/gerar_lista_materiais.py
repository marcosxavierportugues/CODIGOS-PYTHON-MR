# -*- coding: utf-8 -*-
"""
Lista de Materiais IFMT - Gerador Universal
Suporta: qualquer elemento (pilares, vigas, sapatas...)
         qualquer numero de segmentos e pavimentos
         modo RESUMO ou COMPLETO
         nome de arquivo automatico por metadados
"""
import re, os, sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from tkinter import Tk, filedialog, simpledialog, messagebox
import tkinter as tk

# ─── DIALOGS ─────────────────────────────────────────────────────────────────

def _root():
    r = Tk(); r.withdraw(); r.lift(); r.attributes('-topmost', True)
    return r

def _ask_str(title, prompt, default=""):
    r = _root()
    v = simpledialog.askstring(title, prompt, parent=r, initialvalue=default)
    r.destroy()
    return (v or "").strip().upper()

def _ask_file(title):
    r = _root()
    p = filedialog.askopenfilename(title=title, filetypes=[("Excel","*.xlsx *.xls")])
    r.destroy()
    if not p: raise SystemExit(f"Cancelado: {title}")
    return Path(p)

def _ask_save(title, initial_name):
    r = _root()
    p = filedialog.asksaveasfilename(
        title=title, initialfile=initial_name,
        defaultextension=".xlsx", filetypes=[("Excel","*.xlsx")])
    r.destroy()
    if not p: raise SystemExit("Cancelado: salvar arquivo")
    return Path(p)

def _ask_choice(title, prompt, options):
    """Retorna opcao escolhida. options = lista de strings."""
    r = _root()
    var = tk.StringVar(value=options[0])
    win = tk.Toplevel(r)
    win.title(title)
    win.attributes('-topmost', True)
    tk.Label(win, text=prompt, font=("Arial",10), padx=10, pady=8).pack()
    for op in options:
        tk.Radiobutton(win, text=op, variable=var, value=op,
                       font=("Arial",10), anchor="w").pack(fill="x", padx=20)
    result = [options[0]]
    def ok():
        result[0] = var.get(); win.destroy(); r.destroy()
    tk.Button(win, text="OK", command=ok, width=10).pack(pady=8)
    win.wait_window()
    return result[0]

def _ask_yes_no(title, prompt):
    r = _root()
    v = messagebox.askyesno(title, prompt, parent=r)
    r.destroy()
    return v

# ─── METADADOS ───────────────────────────────────────────────────────────────

def coletar_metadados():
    print("\n=== LISTA DE MATERIAIS IFMT ===")
    print("Configurando metadados...\n")

    # Aceitar via env vars (pre-preenchido pelo Claude) ou dialog
    def _env_or_ask_str(env, title, prompt, default=""):
        v = os.environ.get(env, "").strip().upper()
        return v if v else (_ask_str(title, prompt, default) or default.upper())

    elemento  = _env_or_ask_str("LM_ELEMENTO",  "Elemento",  "Qual ELEMENTO?\n(ex: PILARES, VIGAS, SAPATAS, BLOCOS)", "ELEMENTO")
    pavimento = _env_or_ask_str("LM_PAVIMENTO", "Pavimento", "Qual PAVIMENTO?\n(ex: PAV BALDRAMES, TERREO, 1 PAV, 2 PAV)", "PAV")
    segmento  = _env_or_ask_str("LM_SEGMENTO",  "Segmento",  "Qual SEGMENTO?\n(ex: SEG 1, SEG 2, SEG 3, TODOS)", "TODOS")
    obra      = _env_or_ask_str("LM_OBRA",       "Obra",      "Qual OBRA?\n(ex: IFMTR, IFMTC, OBRA XYZ)", "IFMT")

    tipo_env = os.environ.get("LM_TIPO", "").strip().upper()
    tipo = tipo_env if tipo_env in ("EXECUTIVO","OFICIAL") else \
           _ask_choice("Tipo", "Projeto EXECUTIVO ou OFICIAL?", ["EXECUTIVO","OFICIAL"])

    data_env = os.environ.get("LM_DATA", "").strip()
    data_exp = data_env if data_env else \
               (_ask_str("Data", "Data de expedicao?\n(deixar vazio = hoje)",
                         default=datetime.today().strftime("%d.%m.%Y")) or
                datetime.today().strftime("%d.%m.%Y"))

    modo_env = os.environ.get("LM_MODO", "").strip().upper()
    if modo_env in ("A","RESUMO"):
        modo_resumo = True
    elif modo_env in ("B","COMPLETO"):
        modo_resumo = False
    else:
        modo = _ask_choice("Modo de saida",
                           "O que voce quer gerar?",
                           ["(A) So RESUMO (totais por material)",
                            "(B) LISTA COMPLETA (todos os itens + resumo)"])
        modo_resumo = "(A)" in modo

    # Observacao opcional (ex: SEM ESCADAS, BLOCO A, etc.)
    obs_env = os.environ.get("LM_OBS", "").strip().upper()
    if obs_env:
        observacao = obs_env
    else:
        observacao = _ask_str("Observacao",
                              "Observacao opcional (ex: SEM ESCADAS, BLOCO A)\n"
                              "Deixe vazio para nenhuma.", default="").strip().upper()

    nome_base = f"{elemento} - {pavimento} - {obra} - {tipo} - {segmento} - {data_exp}"
    if observacao:
        nome_base += f" - {observacao}"

    print(f"  Elemento  : {elemento}")
    print(f"  Pavimento : {pavimento}")
    print(f"  Segmento  : {segmento}")
    print(f"  Obra      : {obra}")
    print(f"  Tipo      : {tipo}")
    print(f"  Data      : {data_exp}")
    print(f"  Observacao: {observacao or '(nenhuma)'}")
    print(f"  Modo      : {'RESUMO' if modo_resumo else 'COMPLETO'}")
    print(f"  Arquivo   : {nome_base}.xlsx\n")

    return dict(elemento=elemento, pavimento=pavimento, segmento=segmento,
                obra=obra, tipo=tipo, data=data_exp, observacao=observacao,
                nome_base=nome_base, modo_resumo=modo_resumo)

def coletar_arquivos(meta):
    print("Selecione os arquivos de entrada:\n")
    print("1. Planilha TOTAL (todos os elementos)...")
    file_qto = _ask_file("1 - Planilha TOTAL (todos os elementos)")
    print(f"   OK: {file_qto.name}")

    seg_files = {}
    tem_seg = _ask_yes_no("Segmentos", "Voce tem arquivos separados por segmento?\n(SEG 1, SEG 2, etc.)")
    if tem_seg:
        i = 1
        while True:
            continuar = _ask_yes_no("Segmentos", f"Adicionar arquivo SEG {i}?")
            if not continuar:
                break
            print(f"{i+1}. Planilha SEG {i}...")
            f = _ask_file(f"Planilha SEG {i}")
            seg_files[f"SEG {i}"] = f
            print(f"   OK: {f.name}")
            i += 1

    file_out_name = meta['nome_base'].replace("/","-").replace("°","") + ".xlsx"
    file_out = _ask_save("Salvar planilha como...", file_out_name)
    print(f"\nArquivo de saida: {file_out.name}\n")
    return file_qto, seg_files, file_out

# ─── PARSE ───────────────────────────────────────────────────────────────────

NIVEL_MAP = {
    "TERREO":"TERREO","TÉRREO":"TERREO","Térreo":"TERREO","Terreo":"TERREO",
    "1 PAVIMENTO":"TERREO","1 Pavimento":"TERREO","Pavimento 1":"TERREO","PAVIMENTO 1":"TERREO",
    "BALDRAME":"BALDRAME","Baldrame":"BALDRAME","PAV BALDRAMES":"BALDRAME","BALDRAMES":"BALDRAME",
    "2 PAVIMENTO":"PAV 2","2 Pavimento":"PAV 2","Pavimento 2":"PAV 2","2° PAV":"PAV 2",
    "3 PAVIMENTO":"PAV 3","3 Pavimento":"PAV 3","Pavimento 3":"PAV 3","3° PAV":"PAV 3",
}

def normalizar_nivel(v):
    s = str(v).strip()
    return NIVEL_MAP.get(s, s)

def parse_qty(v):
    if pd.isna(v): return 0.0
    txt = str(v).replace(".", "").replace(",", ".")
    m = re.search(r"[\d.]+", txt)
    return float(m.group()) if m else 0.0

def extract_bitola(desc):
    m = re.search(r"[Øø]?\s*([\d,.]+)\s*mm", str(desc), re.IGNORECASE)
    return float(m.group(1).replace(",", ".")) if m else None

def _norm_cols(df):
    """Normaliza nomes de colunas: remove \\r\\n, x000D_, espacos duplos."""
    import re as _re
    rename = {}
    for c in df.columns:
        nc = str(c)
        nc = _re.sub(r"x000D_", "", nc)
        nc = _re.sub(r"[\r\n]+", " ", nc)
        nc = _re.sub(r"\s{2,}", " ", nc).strip()
        if nc != c:
            rename[c] = nc
    if rename:
        df = df.rename(columns=rename)
    return df

def load_qto(path, default_nivel=None):
    """
    Retorna (DataFrame, bitolas, has_layer_col).
    has_layer_col=True  -> arquivo tem coluna Layer com valores reais
    has_layer_col=False -> nivel foi atribuido do metadado (default_nivel)
    """
    df = pd.read_excel(path, sheet_name=0)
    df = _norm_cols(df)
    cols = list(df.columns)
    print(f"  Colunas ({len(cols)}): {cols}")

    # Coluna NOME do elemento
    cand_name = [c for c in cols if "Name" in c and "Layer" not in c]
    if not cand_name:
        cand_name = [c for c in cols if "name" in c.lower() and "layer" not in c.lower()]
    col_name = cand_name[0] if cand_name else cols[0]

    # Coluna NIVEL/LAYER
    cand_layer = [c for c in cols if "Layer" in c
                  and "Layer Id" not in c and "Name" not in c]
    if not cand_layer:
        cand_layer = [c for c in cols if "layer" in c.lower()
                      and "id" not in c.lower() and "name" not in c.lower()]
    col_layer = cand_layer[0] if cand_layer else None
    has_layer_col = col_layer is not None
    if not has_layer_col:
        print(f"  Sem coluna Layer -> nivel = '{default_nivel or 'UNICO'}'")

    # Colunas descricao / quantidade (AltoQi Eberick pattern)
    desc_cols = sorted(c for c in cols if "Descri" in c)
    qty_cols  = sorted(c for c in cols if "Quantidade" in c or "Quantity" in c)

    if not desc_cols or not qty_cols:
        print(f"  ERRO: colunas Descricao/Quantidade nao encontradas: {cols}")
        raise ValueError("Formato de arquivo nao reconhecido. Verifique colunas AltoQi Eberick.")

    print(f"  col_name='{col_name}'  col_layer='{col_layer or '(sem layer)'}'")

    records = []
    for _, row in df.iterrows():
        nivel_raw = row[col_layer] if col_layer else (default_nivel or "UNICO")
        rec = {"ELEM": str(row[col_name]).strip(),
               "NIVEL": normalizar_nivel(nivel_raw),
               "CONCRETO_m3": 0.0, "FORMA_m2": 0.0}
        for dc, qc in zip(desc_cols, qty_cols):
            desc = str(row[dc]) if not pd.isna(row[dc]) else ""
            qty  = parse_qty(row[qc])
            if not desc or qty == 0: continue
            dl = desc.lower()
            if "forma" in dl:
                rec["FORMA_m2"]    += qty
            elif "concreto" in dl:
                rec["CONCRETO_m3"] += qty
            elif any(x in dl for x in ["armadura","ca50","ca60"]):
                b = extract_bitola(desc)
                if b:
                    k = f"ACO_{b}mm_kg"
                    rec[k] = rec.get(k, 0.0) + qty
        records.append(rec)

    result = pd.DataFrame(records)
    aco_cols = sorted((c for c in result.columns if c.startswith("ACO_")),
                      key=lambda x: float(re.search(r"ACO_([\d.]+)mm", x).group(1)))
    bitolas = [float(re.search(r"ACO_([\d.]+)mm", c).group(1))
               for c in aco_cols if result[c].sum() > 0]
    return result, bitolas, has_layer_col

def load_segmentos(seg_files):
    frames = []
    for seg_label, f in seg_files.items():
        df = pd.read_excel(f, sheet_name=0)
        name_cols = [c for c in df.columns if "Name" in c]
        col = name_cols[0] if name_cols else df.columns[1]
        tmp = df[[col]].copy()
        tmp.columns = ["ELEM"]
        tmp["ELEM"]     = tmp["ELEM"].astype(str).str.strip()
        tmp["SEGMENTO"] = seg_label
        frames.append(tmp)
    combined = pd.concat(frames, ignore_index=True)
    return combined.drop_duplicates("ELEM")

# ─── ESTILOS ─────────────────────────────────────────────────────────────────

def fill(c): return PatternFill("solid", start_color=c)

H_PILAR = fill("1F4E79")
H_CONC  = fill("1565C0")
H_FORMA = fill("1E6B3C")
H_ACO   = fill("4A235A")
H_TOTAL = fill("1F4E79")
H_SEM   = fill("C0392B")
S_FILLS = [fill("2E4057"), fill("1B4332"), fill("4A235A"),
           fill("7B241C"), fill("154360"), fill("0B5345")]
ALT     = [fill("F0F4F8"), fill("FFFFFF")]
L_CONC  = fill("D6E4F0")
L_FORMA = fill("D5E8D4")
L_ACO   = fill("EDE7F6")

WF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
NF = Font(name="Arial", size=10)
BF = Font(name="Arial", bold=True, size=10)
C  = Alignment(horizontal="center", vertical="center", wrap_text=True)
L  = Alignment(horizontal="left",   vertical="center")
R  = Alignment(horizontal="right",  vertical="center")
thin = Side(style="thin", color="BFBFBF")
BRD  = Border(left=thin, right=thin, top=thin, bottom=thin)

def _c(ws, r, c, v, fnt=None, fl=None, al=None, fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font      = fnt or NF
    cell.fill      = fl  or PatternFill()
    cell.alignment = al  or R
    cell.border    = BRD
    if fmt: cell.number_format = fmt

def _h(ws, r, c, v, fl=None):
    _c(ws, r, c, v, fnt=WF, fl=fl or H_PILAR, al=C)

def set_widths(ws, extras_cols, bitolas):
    """extras_cols: lista de (nome, largura) antes das cols de qty"""
    widths = [w for _,w in extras_cols] + [13, 13] + [12]*len(bitolas)
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = f"A3"

# ─── WRITE HELPERS ───────────────────────────────────────────────────────────

def write_header(ws, extra_cols, bitolas):
    """extra_cols: [(label, width, fill), ...]"""
    ci = 1
    for lbl, _, fl in extra_cols:
        _h(ws, 1, ci, lbl, fl=fl); _h(ws, 2, ci, "", fl=fl); ci+=1
    _h(ws, 1, ci, "CONCRETO", fl=H_CONC); _h(ws, 2, ci, "m3", fl=H_CONC); ci+=1
    _h(ws, 1, ci, "FORMA",    fl=H_FORMA); _h(ws, 2, ci, "m2", fl=H_FORMA); ci+=1
    for b in bitolas:
        _h(ws, 1, ci, f"O{b}mm", fl=H_ACO); _h(ws, 2, ci, "kg", fl=H_ACO); ci+=1
    ws.freeze_panes = "A3"
    return 3

def write_data_row(ws, ri, row, extra_vals, bitolas, rf):
    ci = 1
    aligns = [L, C, C, C, C, C]
    for idx, val in enumerate(extra_vals):
        _c(ws, ri, ci, val, fnt=NF, fl=rf, al=aligns[idx] if idx<len(aligns) else L); ci+=1
    _c(ws, ri, ci, row.get("CONCRETO_m3",0), fnt=NF, fl=rf, al=R, fmt="#,##0.000"); ci+=1
    _c(ws, ri, ci, row.get("FORMA_m2",0),    fnt=NF, fl=rf, al=R, fmt="#,##0.00");  ci+=1
    for b in bitolas:
        _c(ws, ri, ci, row.get(f"ACO_{b}mm_kg",0), fnt=NF, fl=rf, al=R, fmt="#,##0.000"); ci+=1

def write_total_row(ws, ri, label, sub, bitolas, sfill, n_extra):
    ci = 1
    _c(ws, ri, ci, label, fnt=WF, fl=sfill, al=L); ci+=1
    for _ in range(n_extra-1):
        _c(ws, ri, ci, "", fnt=WF, fl=sfill); ci+=1
    _c(ws, ri, ci, sub["CONCRETO_m3"].sum(), fnt=WF, fl=sfill, al=R, fmt="#,##0.000"); ci+=1
    _c(ws, ri, ci, sub["FORMA_m2"].sum(),    fnt=WF, fl=sfill, al=R, fmt="#,##0.00");  ci+=1
    for b in bitolas:
        col = f"ACO_{b}mm_kg"
        _c(ws, ri, ci, sub[col].sum() if col in sub.columns else 0,
           fnt=WF, fl=sfill, al=R, fmt="#,##0.000"); ci+=1

# ─── ABAS DINÂMICAS ──────────────────────────────────────────────────────────

def build_aba(wb, titulo_aba, titulo_header, sub, bitolas,
              modo_resumo, sfill, grupos=None, label_col=None):
    """
    Cria uma aba.
    grupos: dict {label -> DataFrame} para subtotais dentro da aba
    label_col: nome da coluna a exibir se grupos != None
    """
    ws = wb.create_sheet(titulo_aba)

    if grupos and not modo_resumo:
        # modo completo com subgrupos
        extra_cols = [(titulo_header, 20, H_PILAR)]
        if label_col and label_col != "ELEM":
            extra_cols = [(titulo_header, 20, H_PILAR), (label_col, 12, H_PILAR)]
            if "NIVEL" in sub.columns and label_col != "NIVEL":
                extra_cols += [("NIVEL", 12, H_PILAR)]
        elif "NIVEL" in sub.columns:
            extra_cols += [("NIVEL", 12, H_PILAR)]

        ri = write_header(ws, extra_cols, bitolas)
        for gi, (glabel, gdf) in enumerate(grupos.items()):
            for idx, (_, row) in enumerate(gdf.iterrows()):
                ev = [row.get("ELEM","")]
                if label_col and label_col != "ELEM":
                    ev.append(row.get(label_col,""))
                if "NIVEL" in sub.columns and label_col != "NIVEL":
                    ev.append(row.get("NIVEL",""))
                write_data_row(ws, ri, row, ev, bitolas, ALT[idx%2]); ri+=1
            write_total_row(ws, ri, f"SUBTOTAL {glabel}", gdf, bitolas,
                            S_FILLS[gi % len(S_FILLS)], len(extra_cols)); ri+=1
        write_total_row(ws, ri, "TOTAL GERAL", sub, bitolas, H_TOTAL, len(extra_cols))

    elif not modo_resumo:
        # completo sem subgrupos
        extra_cols = [(titulo_header, 20, H_PILAR)]
        if "NIVEL" in sub.columns:
            extra_cols += [("NIVEL", 12, H_PILAR)]
        ri = write_header(ws, extra_cols, bitolas)
        for idx, (_, row) in enumerate(sub.iterrows()):
            ev = [row.get("ELEM","")]
            if "NIVEL" in sub.columns:
                ev.append(row.get("NIVEL",""))
            write_data_row(ws, ri, row, ev, bitolas, ALT[idx%2]); ri+=1
        write_total_row(ws, ri, "TOTAL", sub, bitolas, H_TOTAL, len(extra_cols))

    else:
        # modo resumo: só totais por subgrupo
        extra_cols = [("DESCRICAO", 30, H_PILAR)]
        ri = write_header(ws, extra_cols, bitolas)
        if grupos:
            for gi, (glabel, gdf) in enumerate(grupos.items()):
                write_total_row(ws, ri, glabel, gdf, bitolas,
                                S_FILLS[gi % len(S_FILLS)], 1); ri+=1
        write_total_row(ws, ri, "TOTAL GERAL", sub, bitolas, H_TOTAL, 1)

    widths = [col[1] for col in (extra_cols if 'extra_cols' in dir() else [(titulo_header,20,None)])]
    for ci, w in enumerate(widths + [13,13] + [12]*len(bitolas), 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    print(f"  Aba '{titulo_aba}': {len(sub)} elementos")
    return ws


def build_total_geral(wb, data, bitolas, meta, grupos_label=None, grupo_col=None):
    """Aba final LISTA DE MATERIAIS / TOTAL GERAL"""
    ws = wb.create_sheet("TOTAL GERAL")

    obs_parte = f" - {meta['observacao']}" if meta.get('observacao') else ""
    titulo = f"LISTA DE MATERIAIS - {meta['elemento']} - {meta['pavimento']} - {meta['obra']} - {meta['tipo']} - {meta['segmento']} - {meta['data']}{obs_parte}"
    ws.merge_cells(f"A1:H1")
    ws["A1"].value     = titulo
    ws["A1"].font      = Font(name="Arial", bold=True, size=12, color="1F4E79")
    ws["A1"].alignment = C
    ws.row_dimensions[1].height = 22

    # cabecalho
    seg_labels = grupos_label or ["TOTAL"]
    hdrs = ["#", "ITEM", "DESCRICAO", "TOTAL GERAL"] + seg_labels + ["UN"]
    hfls = [H_PILAR, H_PILAR, H_PILAR, H_TOTAL] + \
           [S_FILLS[i % len(S_FILLS)] for i in range(len(seg_labels))] + [H_PILAR]
    for ci, (h, f) in enumerate(zip(hdrs, hfls), 1):
        _h(ws, 3, ci, h, fl=f)

    rows_def = [("CONCRETO", "Concreto Estrutural", "CONCRETO_m3", "m3"),
                ("FORMA",    "Forma (Forma)",        "FORMA_m2",    "m2")]
    for b in bitolas:
        rows_def.append(("ACO", f"Armadura CA O {b} mm", f"ACO_{b}mm_kg", "kg"))

    lf = {"CONCRETO": L_CONC, "FORMA": L_FORMA, "ACO": L_ACO}
    ri = 4
    aco_total_row = None

    for idx, (cat, desc, col, un) in enumerate(rows_def, 1):
        fl = lf[cat]
        col_ok = col if col in data.columns else None
        tot = data[col_ok].sum() if col_ok else 0
        # por subgrupo
        if grupos_label and grupo_col and grupo_col in data.columns:
            seg_vals = [data[data[grupo_col] == gl][col_ok].sum()
                        if col_ok else 0 for gl in grupos_label]
        else:
            seg_vals = [tot] * len(seg_labels)

        vals = [idx, cat, desc, tot] + seg_vals + [un]
        for ci, v in enumerate(vals, 1):
            fmt = "#,##0.000" if 4 <= ci <= (4+len(seg_labels)) else None
            al  = L if ci in (2,3) else (R if 4<=ci<=(4+len(seg_labels)) else C)
            _c(ws, ri, ci, v, fl=fl, al=al, fmt=fmt)
        if cat == "ACO" and aco_total_row is None:
            aco_total_row = ri
        ri += 1

    # TOTAL ACO
    aco_cols = [f"ACO_{b}mm_kg" for b in bitolas]
    ws.merge_cells(f"A{ri}:C{ri}")
    _c(ws, ri, 1, "TOTAL ACO", fnt=WF, fl=H_TOTAL, al=C)
    _c(ws, ri, 4, sum(data[c].sum() for c in aco_cols if c in data.columns),
       fnt=WF, fl=H_TOTAL, al=R, fmt="#,##0.000")
    for ci in range(5, 5+len(seg_labels)):
        _c(ws, ri, ci, 0, fnt=WF, fl=H_TOTAL, al=R, fmt="#,##0.000")
    _c(ws, ri, 5+len(seg_labels), "kg", fnt=WF, fl=H_TOTAL, al=C)

    widths = [4, 12, 38, 16] + [14]*len(seg_labels) + [8]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A4"
    print(f"  Aba 'TOTAL GERAL': resumo gerado")

# ─── MAIN ────────────────────────────────────────────────────────────────────

def main():
    meta = coletar_metadados()
    file_qto, seg_files, file_out = coletar_arquivos(meta)

    print("Carregando QTO...")
    data, bitolas, has_layer_col = load_qto(file_qto, default_nivel=meta['pavimento'])
    print(f"  {len(data)} elementos | bitolas: {bitolas}")

    # Se arquivo tem coluna Layer: perguntar qual usar
    if has_layer_col:
        niveis_do_arquivo = sorted(data["NIVEL"].unique())
        print(f"  Niveis no arquivo: {niveis_do_arquivo}")
        if len(niveis_do_arquivo) == 1 and niveis_do_arquivo[0] == meta['pavimento']:
            pass  # Layer col tem mesmo valor que metadado, tudo certo
        else:
            opts = [
                f"Usar Layer do arquivo ({', '.join(str(n) for n in niveis_do_arquivo[:4])}{'...' if len(niveis_do_arquivo)>4 else ''})",
                f"Usar pavimento informado: {meta['pavimento']} (1 nivel unico)",
            ]
            escolha = _ask_choice(
                "Fonte do pavimento",
                f"Arquivo tem coluna Layer com {len(niveis_do_arquivo)} nivel(is).\n"
                f"Qual usar como pavimento?",
                opts
            )
            if "informado" in escolha:
                data["NIVEL"] = meta['pavimento']
                print(f"  Nivel forcado para '{meta['pavimento']}'")
            else:
                print(f"  Usando niveis da coluna Layer do arquivo")

    # atribuir segmentos
    if seg_files:
        print("Carregando segmentos...")
        seg_map = load_segmentos(seg_files)
        data = data.merge(seg_map, on="ELEM", how="left")
        sem = data["SEGMENTO"].isna().sum()
        if sem:
            print(f"  AVISO: {sem} elementos sem segmento -> 'SEM SEGMENTO'")
        data["SEGMENTO"] = data["SEGMENTO"].fillna("SEM SEGMENTO")
    else:
        data["SEGMENTO"] = "TODOS"

    # detectar dimensoes unicas
    segmentos_unicos = [s for s in data["SEGMENTO"].unique() if s != "SEM SEGMENTO"]
    segmentos_unicos.sort()
    niveis_unicos = sorted(data["NIVEL"].unique())
    tem_sem_seg = "SEM SEGMENTO" in data["SEGMENTO"].values

    multi_seg = len(segmentos_unicos) > 1
    multi_niv = len(niveis_unicos) > 1

    print(f"  Segmentos : {segmentos_unicos}")
    print(f"  Niveis    : {niveis_unicos}")

    wb = Workbook()
    wb.remove(wb.active)  # remove aba default

    # ─── Estrutura de abas ─────────────────────────────────────────────────
    if multi_seg and multi_niv:
        # Aba por SEG × NIV
        for si, seg in enumerate(segmentos_unicos):
            for niv in niveis_unicos:
                sub = data[(data["SEGMENTO"]==seg) & (data["NIVEL"]==niv)].reset_index(drop=True)
                if sub.empty: continue
                aba_nome = f"{seg} - {niv}"[:31]
                build_aba(wb, aba_nome, "ELEMENTO", sub, bitolas,
                          meta['modo_resumo'], S_FILLS[si % len(S_FILLS)])
        # Abas por SEG (subtotal)
        for si, seg in enumerate(segmentos_unicos):
            sub = data[data["SEGMENTO"]==seg].reset_index(drop=True)
            grupos = {niv: sub[sub["NIVEL"]==niv].reset_index(drop=True) for niv in niveis_unicos}
            build_aba(wb, seg[:31], "ELEMENTO", sub, bitolas,
                      meta['modo_resumo'], S_FILLS[si % len(S_FILLS)],
                      grupos=grupos, label_col="NIVEL")

    elif multi_seg:
        # Aba por SEG
        for si, seg in enumerate(segmentos_unicos):
            sub = data[data["SEGMENTO"]==seg].reset_index(drop=True)
            build_aba(wb, seg[:31], "ELEMENTO", sub, bitolas,
                      meta['modo_resumo'], S_FILLS[si % len(S_FILLS)])

    elif multi_niv:
        # Aba por NIV
        for niv in niveis_unicos:
            sub = data[data["NIVEL"]==niv].reset_index(drop=True)
            build_aba(wb, niv[:31], "ELEMENTO", sub, bitolas,
                      meta['modo_resumo'], S_FILLS[0])

    else:
        # Aba unica TODOS
        build_aba(wb, "TODOS", "ELEMENTO", data, bitolas,
                  meta['modo_resumo'], H_TOTAL)

    # Aba SEM SEGMENTO
    if tem_sem_seg:
        sub_sem = data[data["SEGMENTO"]=="SEM SEGMENTO"].reset_index(drop=True)
        ws_sem = build_aba(wb, "SEM SEGMENTO", "ELEMENTO", sub_sem, bitolas,
                           meta['modo_resumo'], H_SEM)
        print(f"  ATENCAO: {len(sub_sem)} elementos em 'SEM SEGMENTO'")

    # Aba TOTAL GERAL sempre no final
    grupos_label = segmentos_unicos if multi_seg else (niveis_unicos if multi_niv else None)
    grupo_col    = "SEGMENTO"        if multi_seg else ("NIVEL"       if multi_niv else None)
    build_total_geral(wb, data, bitolas, meta, grupos_label, grupo_col)

    wb.save(file_out)
    # copiar para D:\Cache\Claude\
    try:
        import shutil
        cache = Path(r"D:\Cache\Claude")
        cache.mkdir(parents=True, exist_ok=True)
        shutil.copy2(file_out, cache / file_out.name)
        print(f"\nCopiado para: D:\\Cache\\Claude\\{file_out.name}")
    except Exception as e:
        print(f"  (copia cache falhou: {e})")

    print(f"\nSalvo em:\n{file_out}")
    print(f"\nTotal : {len(data)} elementos")
    print(f"Bitolas: {bitolas}")
    for seg in segmentos_unicos:
        n = (data["SEGMENTO"]==seg).sum()
        print(f"  {seg}: {n} elementos")
    for niv in niveis_unicos:
        n = (data["NIVEL"]==niv).sum()
        print(f"  {niv}: {n} elementos")

    # Perguntar se continua ou finaliza
    continuar = _ask_yes_no(
        "Concluido",
        f"Planilha gerada:\n{file_out.name}\n\n"
        "Deseja gerar uma NOVA planilha?\n\n"
        "SIM = novo processo     NAO = finalizar"
    )
    return continuar

if __name__ == "__main__":
    # Limpar env vars de sessao anterior para forcara dialogs na proxima rodada
    ENV_VARS = ["LM_ELEMENTO","LM_PAVIMENTO","LM_SEGMENTO","LM_OBRA",
                "LM_TIPO","LM_DATA","LM_MODO","LM_OBS"]
    while True:
        continuar = main()
        # Apos primeira rodada, limpar env vars para que proxima rodada
        # use os dialogs (usuario pode querer dados diferentes)
        for v in ENV_VARS:
            os.environ.pop(v, None)
        if not continuar:
            print("\nFinalizado. Ate logo!")
            break
