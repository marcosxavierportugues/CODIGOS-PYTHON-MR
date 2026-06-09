import pandas as pd
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

file_in  = r'D:\00-Obras - SLN\IFMT - SLN\00- ENVIADO IFMT\SEG 1\LSEG 1- LAJES FAIXA.xlsx'
file_out = r'D:\00-Obras - SLN\IFMT - SLN\00- ENVIADO IFMT\SEG 1\LSEG 1- LAJES FAIXA - ORGANIZADO-BITOLA.xlsx'

df = pd.read_excel(file_in, sheet_name='Sheet1', header=0)
df.columns = [str(c) for c in df.columns]

COL_NAME  = 0
COL_LAYER = 6
item_pairs = [(58,59),(60,61),(62,63),(64,65),(66,67),(68,69),(70,71),(72,73),(74,75),(76,77),(78,79)]

# Special non-kg items (punção) — tracked by key
PUNCAO_CA25  = 'PUNCAO_CA25_O8mm'
CHAPA_RETANG = 'CHAPA_RETANG_16.2x3.2'


def parse_qty(val):
    """Extract numeric value from '1.388 m³', '192.49 kg', '15.0', etc."""
    if pd.isna(val):
        return None
    m = re.search(r'([\d\.]+)', str(val))
    return float(m.group(1)) if m else None


def extract_bitola(desc):
    # Match first decimal number followed by mm
    m = re.search(r'([\d\.]+)\s*mm', str(desc))
    return float(m.group(1)) if m else None


def classify(desc):
    d = str(desc)
    if 'Forma' in d:
        return 'FORMA'
    if 'Concreto' in d:
        return 'CONCRETO'
    if 'Chapa retangular' in d or 'Chapa' in d:
        return 'CHAPA'
    if 'pun' in d.lower():
        return 'PUNCAO'
    if 'Armadura' in d or 'rmadura' in d:
        return 'ACO'
    return None


# Collect regular bitolas (CA50 type)
all_bitolas = set()
for _, row in df.iterrows():
    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        if pd.notna(desc) and classify(desc) == 'ACO':
            b = extract_bitola(desc)
            if b:
                all_bitolas.add(b)

bitolas = sorted(all_bitolas)
print('Bitolas CA50 encontradas:', bitolas)

# Check if punção items exist
has_puncao = False
has_chapa  = False
for _, row in df.iterrows():
    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        if pd.notna(desc):
            cat = classify(desc)
            if cat == 'PUNCAO':
                has_puncao = True
            if cat == 'CHAPA':
                has_chapa = True

print(f'Tem punção CA25: {has_puncao} | Tem chapa retangular: {has_chapa}')

# Build structured records
records = []
for _, row in df.iterrows():
    elem  = str(row.iloc[COL_NAME])
    layer = str(row.iloc[COL_LAYER])
    concreto = 0.0
    forma = 0.0
    aco = {b: 0.0 for b in bitolas}
    puncao = 0.0
    chapa  = 0.0

    for dc, qc in item_pairs:
        desc = row.iloc[dc]
        qty_raw = row.iloc[qc]
        if pd.isna(desc):
            continue
        cat = classify(desc)
        qty = parse_qty(qty_raw)
        if qty is None:
            continue
        if cat == 'CONCRETO':
            concreto += qty
        elif cat == 'FORMA':
            forma += qty
        elif cat == 'ACO':
            b = extract_bitola(desc)
            if b and b in aco:
                aco[b] += qty
        elif cat == 'PUNCAO':
            puncao += qty
        elif cat == 'CHAPA':
            chapa += qty

    rec = {'LAJE': elem, 'CAMADA': layer, 'CONCRETO_m3': concreto, 'FORMA_m2': forma}
    for b in bitolas:
        rec[f'ACO_{b}mm_kg'] = aco[b]
    if has_puncao:
        rec['PUNCAO_CA25_un'] = puncao
    if has_chapa:
        rec['CHAPA_RETANG_un'] = chapa
    records.append(rec)

result = pd.DataFrame(records)

# VALIDATION
print('\n=== VALIDACAO - TOTAIS GERAIS ===')
print(f'Total CONCRETO: {result["CONCRETO_m3"].sum():.4f} m3')
print(f'Total FORMA:    {result["FORMA_m2"].sum():.4f} m2')
for b in bitolas:
    t = result[f'ACO_{b}mm_kg'].sum()
    if t > 0:
        print(f'Total ACO O{b}mm: {t:.4f} kg')
if has_puncao:
    print(f'Total PUNCAO CA25 O8mm: {result["PUNCAO_CA25_un"].sum():.0f} un')
if has_chapa:
    print(f'Total CHAPA RETANG: {result["CHAPA_RETANG_un"].sum():.0f} un')
print(f'Total lajes: {len(result)}')

# ========================
# BUILD EXCEL
# ========================
wb = openpyxl.Workbook()

thin = Side(style='thin', color='BBBBBB')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_al = Alignment(horizontal='left', vertical='center')
right_al = Alignment(horizontal='right', vertical='center')

white_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
normal_font = Font(name='Arial', size=10)

hdr_blue   = PatternFill('solid', start_color='1F4E79')
hdr_green  = PatternFill('solid', start_color='375623')
hdr_purple = PatternFill('solid', start_color='4A235A')
hdr_red    = PatternFill('solid', start_color='7B1818')
hdr_brown  = PatternFill('solid', start_color='5C3317')
fill_total = PatternFill('solid', start_color='1F4E79')
fill_gold  = PatternFill('solid', start_color='D4AC0D')
fill_conc  = PatternFill('solid', start_color='D6E4F0')
fill_forma = PatternFill('solid', start_color='D5E8D4')
fill_aco   = PatternFill('solid', start_color='EDE7F6')
fill_punc  = PatternFill('solid', start_color='FDEBD0')

# =============================================
# ABA 1 - DETALHADO
# =============================================
ws = wb.active
ws.title = 'DETALHADO'

header_labels = ['LAJE', 'CAMADA', 'CONCRETO', 'FORMA'] + [f'ACO O{b}mm' for b in bitolas]
unit_labels   = ['', '', 'm3', 'm2'] + ['kg'] * len(bitolas)
fills_hdr     = [hdr_blue, hdr_blue, hdr_blue, hdr_green] + [hdr_purple] * len(bitolas)

if has_puncao:
    header_labels.append('PUNCAO CA25 O8mm')
    unit_labels.append('un')
    fills_hdr.append(hdr_red)
if has_chapa:
    header_labels.append('CHAPA RETANG 16.2x3.2')
    unit_labels.append('un')
    fills_hdr.append(hdr_brown)

for ci, (h, u, f) in enumerate(zip(header_labels, unit_labels, fills_hdr), 1):
    for ri, val in [(1, h), (2, u)]:
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = white_font; c.fill = f; c.alignment = center; c.border = border

for ri, rec in enumerate(records, 3):
    row_fill = PatternFill('solid', start_color='F5F5F5') if ri % 2 == 1 else PatternFill('solid', start_color='FFFFFF')
    vals = [rec['LAJE'], rec['CAMADA'], rec['CONCRETO_m3'], rec['FORMA_m2']] + \
           [rec[f'ACO_{b}mm_kg'] for b in bitolas]
    if has_puncao:
        vals.append(rec['PUNCAO_CA25_un'])
    if has_chapa:
        vals.append(rec['CHAPA_RETANG_un'])
    for ci, val in enumerate(vals, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.font = normal_font; c.border = border; c.fill = row_fill
        c.alignment = left_al if ci <= 2 else right_al
        if ci > 2:
            c.number_format = '#,##0.00'

ri_tot = 3 + len(records)
tot_vals = ['TOTAL', '', result['CONCRETO_m3'].sum(), result['FORMA_m2'].sum()] + \
           [result[f'ACO_{b}mm_kg'].sum() for b in bitolas]
if has_puncao:
    tot_vals.append(result['PUNCAO_CA25_un'].sum())
if has_chapa:
    tot_vals.append(result['CHAPA_RETANG_un'].sum())
for ci, v in enumerate(tot_vals, 1):
    c = ws.cell(row=ri_tot, column=ci, value=v)
    c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    c.fill = fill_total; c.border = border
    c.alignment = left_al if ci <= 2 else right_al
    if ci > 2 and isinstance(v, float):
        c.number_format = '#,##0.00'

col_widths = [14, 20, 14, 14] + [13] * len(bitolas) + ([16, 22] if has_puncao and has_chapa else [])
for ci, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.row_dimensions[1].height = 35
ws.row_dimensions[2].height = 18
ws.freeze_panes = 'A3'

# =============================================
# ABA 2 - RESUMO TOTAL
# =============================================
ws2 = wb.create_sheet('RESUMO TOTAL')
ws2['A1'] = 'RESUMO GERAL - SEG 1 LAJES FAIXA'
ws2['A1'].font = Font(name='Arial', bold=True, size=14, color='1F4E79')
ws2.merge_cells('A1:D1')
ws2['A1'].alignment = center

for ci, h in enumerate(['CATEGORIA', 'DESCRICAO', 'TOTAL', 'UNIDADE'], 1):
    c = ws2.cell(row=3, column=ci, value=h)
    c.font = white_font; c.fill = hdr_blue; c.alignment = center; c.border = border

rows_resumo = [('CONCRETO', 'Concreto C-30 - Abatimento 5 cm', result['CONCRETO_m3'].sum(), 'm3'),
               ('FORMA', 'Forma - Estrutura - Concreto', result['FORMA_m2'].sum(), 'm2')]
for b in bitolas:
    t = result[f'ACO_{b}mm_kg'].sum()
    if t > 0:
        rows_resumo.append(('ACO', f'Armadura CA50 - O {b} mm', t, 'kg'))
if has_puncao:
    rows_resumo.append(('PUNCAO', 'Armadura de puncao - Aco CA25 - O 8.0 mm (9 cm)',
                        result['PUNCAO_CA25_un'].sum(), 'un'))
if has_chapa:
    rows_resumo.append(('CHAPA', 'Armadura de puncao - Chapa retangular 16.2x3.2',
                        result['CHAPA_RETANG_un'].sum(), 'un'))

fills_cat = {'CONCRETO': fill_conc, 'FORMA': fill_forma, 'ACO': fill_aco,
             'PUNCAO': fill_punc, 'CHAPA': fill_punc}
for ri, (cat, desc, val, unit) in enumerate(rows_resumo, 4):
    f = fills_cat.get(cat, fill_conc)
    for ci, v in enumerate([cat, desc, val, unit], 1):
        c = ws2.cell(row=ri, column=ci, value=v)
        c.font = normal_font; c.border = border; c.fill = f
        c.alignment = left_al if ci == 2 else center
        if ci == 3:
            c.number_format = '#,##0.00'

r_gt = 4 + len(rows_resumo)
aco_grand = sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
for ci, v in enumerate(['ACO TOTAL', 'Total Geral de Aco CA50 (todas as bitolas)', aco_grand, 'kg'], 1):
    c = ws2.cell(row=r_gt, column=ci, value=v)
    c.font = Font(name='Arial', bold=True, size=10); c.fill = fill_gold; c.border = border
    c.alignment = left_al if ci == 2 else center
    if ci == 3:
        c.number_format = '#,##0.00'

ws2.column_dimensions['A'].width = 14
ws2.column_dimensions['B'].width = 50
ws2.column_dimensions['C'].width = 16
ws2.column_dimensions['D'].width = 12

# =============================================
# ABA 3 - RESUMO POR CAMADA
# =============================================
ws3 = wb.create_sheet('RESUMO POR CAMADA')
extra_cols = (1 if has_puncao else 0) + (1 if has_chapa else 0)
n_cols_3 = 3 + len(bitolas) + 1 + extra_cols
ws3['A1'] = 'RESUMO POR CAMADA - SEG 1 LAJES FAIXA'
ws3['A1'].font = Font(name='Arial', bold=True, size=14, color='1F4E79')
ws3.merge_cells(f'A1:{get_column_letter(n_cols_3)}1')
ws3['A1'].alignment = center

hdrs3 = ['CAMADA', 'CONCRETO (m3)', 'FORMA (m2)'] + [f'ACO O{b}mm (kg)' for b in bitolas] + ['ACO TOTAL (kg)']
if has_puncao:
    hdrs3.append('PUNCAO CA25 (un)')
if has_chapa:
    hdrs3.append('CHAPA RETANG (un)')

for ci, h in enumerate(hdrs3, 1):
    c = ws3.cell(row=3, column=ci, value=h)
    c.font = white_font; c.fill = hdr_blue; c.alignment = center; c.border = border

camadas = result['CAMADA'].unique()
for ri, cam in enumerate(camadas, 4):
    sub = result[result['CAMADA'] == cam]
    aco_tot = sum(sub[f'ACO_{b}mm_kg'].sum() for b in bitolas)
    vals3 = [cam, sub['CONCRETO_m3'].sum(), sub['FORMA_m2'].sum()] + \
            [sub[f'ACO_{b}mm_kg'].sum() for b in bitolas] + [aco_tot]
    if has_puncao:
        vals3.append(sub['PUNCAO_CA25_un'].sum())
    if has_chapa:
        vals3.append(sub['CHAPA_RETANG_un'].sum())
    row_fill = PatternFill('solid', start_color='EBF3FB') if ri % 2 == 0 else PatternFill('solid', start_color='FFFFFF')
    for ci, v in enumerate(vals3, 1):
        c = ws3.cell(row=ri, column=ci, value=v)
        c.font = normal_font; c.border = border; c.fill = row_fill
        c.alignment = left_al if ci == 1 else right_al
        if ci > 1:
            c.number_format = '#,##0.00'

ri_tot3 = 4 + len(camadas)
aco_all = sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
totals3 = ['TOTAL GERAL', result['CONCRETO_m3'].sum(), result['FORMA_m2'].sum()] + \
          [result[f'ACO_{b}mm_kg'].sum() for b in bitolas] + [aco_all]
if has_puncao:
    totals3.append(result['PUNCAO_CA25_un'].sum())
if has_chapa:
    totals3.append(result['CHAPA_RETANG_un'].sum())
for ci, v in enumerate(totals3, 1):
    c = ws3.cell(row=ri_tot3, column=ci, value=v)
    c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    c.fill = fill_total; c.border = border
    c.alignment = left_al if ci == 1 else right_al
    if ci > 1:
        c.number_format = '#,##0.00'

for ci, w in enumerate([22, 16, 16] + [13] * len(bitolas) + [16] + [18]*extra_cols, 1):
    ws3.column_dimensions[get_column_letter(ci)].width = w

wb.save(file_out)
print(f'\nArquivo salvo: {file_out}')
