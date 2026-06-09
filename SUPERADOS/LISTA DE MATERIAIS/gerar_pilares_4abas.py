"""
Gera planilha de pilares com 5 abas:
  1. TODOS OS PILARES
  2. SEG 1
  3. SEG 2
  4. SEG 3
  5. LISTA DE MATERIAIS
"""
import re
import pandas as pd
from pathlib import Path
from tkinter import Tk, filedialog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# SELEÇÃO DE ARQUIVOS
# ─────────────────────────────────────────────
def escolher(titulo, salvar=False):
    root = Tk(); root.withdraw()
    if salvar:
        p = filedialog.asksaveasfilename(
            title=titulo, defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")])
    else:
        p = filedialog.askopenfilename(
            title=titulo,
            filetypes=[("Excel", "*.xlsx *.xls")])
    root.destroy()
    if not p:
        raise SystemExit(f"Nenhum arquivo selecionado: {titulo}")
    return Path(p)

FILE_QTO  = escolher("Selecione: TODOS OS PILARES (QTO)")
SEG_FILES = [
    escolher("Selecione: PILARES SEG 1"),
    escolher("Selecione: PILARES SEG 2"),
    escolher("Selecione: PILARES SEG 3"),
]
FILE_OUT  = escolher("Salvar planilha como", salvar=True)

TITULO    = "IFMT - PILARES TÉRREO"
SEGMENTOS = ["SEG 1", "SEG 2", "SEG 3"]

# ─────────────────────────────────────────────
# ESTILOS
# ─────────────────────────────────────────────
def fill(c):    return PatternFill("solid", start_color=c)
def font(bold=False, color="000000", size=10):
    return Font(name="Arial", bold=bold, color=color, size=size)

HDR_PILAR  = fill("1F4E79")   # azul escuro
HDR_CONC   = fill("1565C0")   # azul médio
HDR_FORMA  = fill("1E6B3C")   # verde
HDR_ACO    = fill("4A235A")   # roxo
ROW_ALT    = [fill("F0F4F8"), fill("FFFFFF")]
SUB_FILLS  = [fill("2E4057"), fill("1B4332"), fill("4A235A")]   # subtotal por seg
TOTAL_FILL = fill("1F4E79")
LISTA_FILLS = {"CONCRETO": fill("D6E4F0"),
               "FORMA":    fill("D5E8D4"),
               "ACO":      fill("EDE7F6")}

WF    = font(bold=True, color="FFFFFF")
NF    = font()
BF    = font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center")
RIGHT  = Alignment(horizontal="right",  vertical="center")
thin   = Side(style="thin", color="BFBFBF")
BRD    = Border(left=thin, right=thin, top=thin, bottom=thin)

def _c(ws, r, c, val, fnt=None, fl=None, al=None, fmt=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font      = fnt or NF
    cell.fill      = fl  or PatternFill()
    cell.alignment = al  or RIGHT
    cell.border    = BRD
    if fmt: cell.number_format = fmt

def _h(ws, r, c, val, fl=None):
    _c(ws, r, c, val, fnt=WF, fl=fl or HDR_PILAR, al=CENTER)

# ─────────────────────────────────────────────
# HELPERS DE PARSE
# ─────────────────────────────────────────────
def parse_qty(val):
    if pd.isna(val): return 0.0
    txt = str(val).replace(".", "").replace(",", ".")
    m = re.search(r"[\d.]+", txt)
    return float(m.group()) if m else 0.0

def extract_bitola(desc):
    m = re.search(r"[Øø⌀-]?\s*([\d,.]+)\s*mm", str(desc), re.IGNORECASE)
    return float(m.group(1).replace(",", ".")) if m else None

def corrigir_nivel(v):
    mapa = {"1 Pavimento":"TERREO","1 PAVIMENTO":"TERREO","Térreo":"TERREO",
            "TÉRREO":"TERREO","Terreo":"TERREO","TERREO":"TERREO",
            "Pavimento 1":"TERREO","PAVIMENTO 1":"TERREO"}
    return mapa.get(str(v).strip(), str(v).strip())

# ─────────────────────────────────────────────
# CARGA DE DADOS
# ─────────────────────────────────────────────
def load_qto(path):
    df = pd.read_excel(path, sheet_name=0)
    col_name  = [c for c in df.columns if "Name"  in c][0]
    col_layer = [c for c in df.columns
                 if "Layer" in c and "Layer Id" not in c and "Name" not in c][0]
    desc_cols = sorted([c for c in df.columns if "Descri"    in c])
    qty_cols  = sorted([c for c in df.columns if "Quantidade" in c])

    records = []
    for _, row in df.iterrows():
        rec = {"PILAR": str(row[col_name]).strip(),
               "NIVEL": corrigir_nivel(row[col_layer]),
               "CONCRETO_m3": 0.0, "FORMA_m2": 0.0}
        for dc, qc in zip(desc_cols, qty_cols):
            desc = str(row[dc]) if not pd.isna(row[dc]) else ""
            qty  = parse_qty(row[qc])
            if not desc or qty == 0: continue
            if "Forma"   in desc: rec["FORMA_m2"]    += qty
            elif "Concreto" in desc: rec["CONCRETO_m3"] += qty
            elif any(x in desc for x in ["Armadura","CA50","CA60"]):
                b = extract_bitola(desc)
                if b:
                    k = f"ACO_{b}mm_kg"
                    rec[k] = rec.get(k, 0.0) + qty
        records.append(rec)

    result = pd.DataFrame(records)
    aco_cols = sorted([c for c in result.columns if c.startswith("ACO_")],
                      key=lambda x: float(re.search(r"ACO_([\d.]+)mm", x).group(1)))
    bitolas = [float(re.search(r"ACO_([\d.]+)mm", c).group(1))
               for c in aco_cols if result[c].sum() > 0]
    for b in bitolas:
        k = f"ACO_{b}mm_kg"
        if k not in result.columns:
            result[k] = 0.0
        result[k] = result[k].fillna(0.0)
    return result, bitolas

def load_segmentos(seg_files):
    frames = []
    for i, f in enumerate(seg_files, 1):
        df = pd.read_excel(f, sheet_name=0)
        tmp = df[[df.columns[0]]].copy()
        tmp.columns = ["PILAR"]
        tmp["PILAR"]    = tmp["PILAR"].astype(str).str.strip()
        tmp["SEGMENTO"] = f"SEG {i}"
        frames.append(tmp)
    return pd.concat(frames, ignore_index=True).drop_duplicates("PILAR")

# ─────────────────────────────────────────────
# CONSTRUÇÃO DE ABA (reutilizável)
# ─────────────────────────────────────────────
def col_widths(bitolas):
    return [16, 10, 12, 14, 14] + [13]*len(bitolas)

def write_header(ws, bitolas, show_seg=True):
    """Escreve 2 linhas de cabeçalho. Retorna próxima linha."""
    cols  = ["PILAR"]
    hfill = [HDR_PILAR]
    if show_seg:
        cols  += ["SEGMENTO"]
        hfill += [HDR_PILAR]
    cols  += ["NÍVEL",   "CONCRETO",  "FORMA"]
    hfill += [HDR_PILAR, HDR_CONC, HDR_FORMA]
    cols  += [f"ACO Ø{b}mm" for b in bitolas]
    hfill += [HDR_ACO] * len(bitolas)

    units = ["", ""] if show_seg else [""]
    units += ["", "m³", "m²"] + ["kg"]*len(bitolas)

    for ci, (h, f) in enumerate(zip(cols, hfill), 1):
        _h(ws, 1, ci, h, fl=f)
        _h(ws, 2, ci, units[ci-1], fl=f)
    return 3

def write_pilar_row(ws, ri, row, bitolas, seg=None, row_fill=None):
    rf = row_fill or PatternFill()
    ci = 1
    _c(ws, ri, ci, row["PILAR"],       fnt=NF, fl=rf, al=LEFT);  ci+=1
    if seg is not None:
        _c(ws, ri, ci, seg,            fnt=NF, fl=rf, al=CENTER); ci+=1
    _c(ws, ri, ci, row["NIVEL"],       fnt=NF, fl=rf, al=LEFT);   ci+=1
    _c(ws, ri, ci, row["CONCRETO_m3"], fnt=NF, fl=rf, al=RIGHT, fmt="#,##0.00"); ci+=1
    _c(ws, ri, ci, row["FORMA_m2"],    fnt=NF, fl=rf, al=RIGHT, fmt="#,##0.00"); ci+=1
    for b in bitolas:
        _c(ws, ri, ci, row[f"ACO_{b}mm_kg"], fnt=NF, fl=rf, al=RIGHT, fmt="#,##0.00"); ci+=1

def write_subtotal(ws, ri, label, sub, bitolas, sfill, show_seg=True):
    ci = 1
    _c(ws, ri, ci, label, fnt=WF, fl=sfill, al=LEFT); ci+=1
    if show_seg:
        _c(ws, ri, ci, "", fnt=WF, fl=sfill); ci+=1
    _c(ws, ri, ci, "", fnt=WF, fl=sfill); ci+=1
    _c(ws, ri, ci, sub["CONCRETO_m3"].sum(), fnt=WF, fl=sfill, al=RIGHT, fmt="#,##0.00"); ci+=1
    _c(ws, ri, ci, sub["FORMA_m2"].sum(),    fnt=WF, fl=sfill, al=RIGHT, fmt="#,##0.00"); ci+=1
    for b in bitolas:
        _c(ws, ri, ci, sub[f"ACO_{b}mm_kg"].sum(), fnt=WF, fl=sfill, al=RIGHT, fmt="#,##0.00"); ci+=1

def write_total(ws, ri, data, bitolas, show_seg=True):
    write_subtotal(ws, ri, "TOTAL GERAL", data, bitolas, TOTAL_FILL, show_seg=show_seg)

def apply_col_widths(ws, bitolas, show_seg=True):
    widths = [16]
    if show_seg: widths.append(10)
    widths += [12, 14, 14] + [13]*len(bitolas)
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

# ─────────────────────────────────────────────
# ABA 1 — TODOS OS PILARES
# ─────────────────────────────────────────────
def build_todos(wb, data, bitolas):
    ws = wb.active
    ws.title = "TODOS OS PILARES"
    ri = write_header(ws, bitolas, show_seg=True)

    for si, seg in enumerate(SEGMENTOS):
        sub = data[data["SEGMENTO"] == seg]
        for idx, (_, row) in enumerate(sub.iterrows()):
            write_pilar_row(ws, ri, row, bitolas, seg=seg, row_fill=ROW_ALT[idx%2])
            ri += 1
        write_subtotal(ws, ri, f"SUBTOTAL {seg}", sub, bitolas, SUB_FILLS[si], show_seg=True)
        ri += 1

    write_total(ws, ri, data, bitolas, show_seg=True)
    apply_col_widths(ws, bitolas, show_seg=True)
    ws.freeze_panes = "A3"

# ─────────────────────────────────────────────
# ABA 2/3/4 — POR SEGMENTO
# ─────────────────────────────────────────────
def build_segmento(wb, data, bitolas, seg, sfill):
    ws = wb.create_sheet(seg)
    sub = data[data["SEGMENTO"] == seg].reset_index(drop=True)
    ri  = write_header(ws, bitolas, show_seg=False)

    for idx, (_, row) in enumerate(sub.iterrows()):
        write_pilar_row(ws, ri, row, bitolas, seg=None, row_fill=ROW_ALT[idx%2])
        ri += 1

    write_subtotal(ws, ri, f"TOTAL {seg}", sub, bitolas, sfill, show_seg=False)
    apply_col_widths(ws, bitolas, show_seg=False)
    ws.freeze_panes = "A3"
    print(f"  {seg}: {len(sub)} pilares")

# ─────────────────────────────────────────────
# ABA 5 — LISTA DE MATERIAIS
# ─────────────────────────────────────────────
def build_lista_materiais(wb, data, bitolas):
    ws = wb.create_sheet("LISTA DE MATERIAIS")

    # Título
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value     = f"LISTA DE MATERIAIS — {TITULO}"
    c.font      = Font(name="Arial", bold=True, size=14, color="1F4E79")
    c.alignment = CENTER
    ws.row_dimensions[1].height = 22

    # Cabeçalho
    hdrs = ["#", "ITEM", "DESCRIÇÃO", "TODOS", "SEG 1", "SEG 2", "SEG 3", "UNIDADE"]
    fills_hdr = [HDR_PILAR]*3 + [TOTAL_FILL] + SUB_FILLS + [HDR_PILAR]
    for ci, (h, f) in enumerate(zip(hdrs, fills_hdr), 1):
        _h(ws, 3, ci, h, fl=f)

    # Linhas de materiais
    items = []
    items.append(("CONCRETO", "Concreto Estrutural", "CONCRETO", "m³"))
    items.append(("FORMA",    "Forma (Fôrma)",       "FORMA",    "m²"))
    for b in bitolas:
        items.append(("ACO", f"Armadura CA Ø {b} mm", f"ACO_{b}", "kg"))

    ri = 4
    idx_num = 1
    for (cat, desc, key, unit) in items:
        fl = LISTA_FILLS.get(cat, fill("FFFFFF"))

        if cat == "CONCRETO":
            tot   = data["CONCRETO_m3"].sum()
            s_vals = [data[data["SEGMENTO"]==s]["CONCRETO_m3"].sum() for s in SEGMENTOS]
        elif cat == "FORMA":
            tot   = data["FORMA_m2"].sum()
            s_vals = [data[data["SEGMENTO"]==s]["FORMA_m2"].sum() for s in SEGMENTOS]
        else:
            col   = f"ACO_{key.split('_')[1]}mm_kg"
            tot   = data[col].sum()
            s_vals = [data[data["SEGMENTO"]==s][col].sum() for s in SEGMENTOS]

        row_vals = [idx_num, cat, desc, tot] + s_vals + [unit]
        for ci, v in enumerate(row_vals, 1):
            fmt = "#,##0.00" if 4 <= ci <= 7 else None
            al  = LEFT if ci in (2,3) else (RIGHT if 4<=ci<=7 else CENTER)
            _c(ws, ri, ci, v, fnt=(BF if ci==3 else NF), fl=fl, al=al, fmt=fmt)
        ri += 1
        idx_num += 1

    # Total de aço
    ws.merge_cells(f"A{ri}:C{ri}")
    _c(ws, ri, 1, "TOTAL AÇO", fnt=WF, fl=TOTAL_FILL, al=CENTER)
    aco_cols = [f"ACO_{b}mm_kg" for b in bitolas]
    tot_aco = data[aco_cols].sum().sum()
    s_aco   = [data[data["SEGMENTO"]==s][aco_cols].sum().sum() for s in SEGMENTOS]
    for ci, v in enumerate([tot_aco]+s_aco, 4):
        _c(ws, ri, ci, v, fnt=WF, fl=TOTAL_FILL, al=RIGHT, fmt="#,##0.00")
    _c(ws, ri, 8, "kg", fnt=WF, fl=TOTAL_FILL, al=CENTER)

    # Larguras
    for ci, w in enumerate([5, 12, 38, 16, 14, 14, 14, 10], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A4"

# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────
def main():
    print("Carregando QTO...")
    data, bitolas = load_qto(FILE_QTO)

    print("Carregando segmentos...")
    seg_map = load_segmentos(SEG_FILES)
    data = data.merge(seg_map, on="PILAR", how="left")
    sem = data["SEGMENTO"].isna().sum()
    if sem:
        print(f"  ATENÇÃO: {sem} pilares sem segmento → classificados como 'SEM SEG'")
        data["SEGMENTO"] = data["SEGMENTO"].fillna("SEM SEG")

    ord_map = {s: i for i, s in enumerate(SEGMENTOS)}
    data["_ord"] = data["SEGMENTO"].map(ord_map).fillna(99)
    data = data.sort_values(["_ord", "PILAR"]).drop(columns="_ord").reset_index(drop=True)

    print(f"\nTotal de pilares : {len(data)}")
    print(f"Bitolas          : {bitolas}")
    print(f"Por segmento:")
    for s in SEGMENTOS:
        print(f"  {s}: {(data['SEGMENTO']==s).sum()} pilares")

    print("\nGerando abas...")
    wb = Workbook()
    build_todos(wb, data, bitolas)
    for i, (seg, sfill) in enumerate(zip(SEGMENTOS, SUB_FILLS)):
        build_segmento(wb, data, bitolas, seg, sfill)
    build_lista_materiais(wb, data, bitolas)

    wb.save(FILE_OUT)
    print(f"\nSalvo em: {FILE_OUT}")

if __name__ == "__main__":
    main()
