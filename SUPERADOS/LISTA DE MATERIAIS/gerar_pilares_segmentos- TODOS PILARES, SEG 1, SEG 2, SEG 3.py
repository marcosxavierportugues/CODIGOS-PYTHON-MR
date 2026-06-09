"""
Gerador de planilha de PILARES TÉRREO organizada por SEGMENTO.
Fontes:
  - IFMT TODOS OS PILARES TERREO.xlsx  (quantitativos)
  - PILARES TERREO/IFMT - ENVIADO  PILARES TERREO SEG 1/2/3.xlsx (mapeamento segmento)
Abas geradas: DETALHADO | RESUMO TOTAL | RESUMO POR SEGMENTO | PRORRATA
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── CONFIG ───────────────────────────────────────────────────────────────────
BASE     = Path(r"D:\00-Obras - SLN\IFMT - SLN\0000- NAVISWORKS IFMT")
FILE_QTO = BASE / "IFMT TODOS OS PILARES TERREO.xlsx"
SEG_FILES = [
    BASE / "PILARES TERREO" / "IFMT - ENVIADO  PILARES TERREO SEG 1.xlsx",
    BASE / "PILARES TERREO" / "IFMT - ENVIADO  PILARES TERREO SEG 2.xlsx",
    BASE / "PILARES TERREO" / "IFMT - ENVIADO  PILARES TERREO SEG 3.xlsx",
]
FILE_OUT = BASE / "IFMT TODOS OS PILARES TERREO_SEGMENTOS.xlsx"
TITULO   = "IFMT - PILARES TÉRREO"
SEGMENTOS = ["SEG 1", "SEG 2", "SEG 3"]

# ── ESTILOS ──────────────────────────────────────────────────────────────────
def fill(c): return PatternFill('solid', start_color=c)

hdr_blue   = fill('1F4E79'); hdr_green = fill('1E6B3C'); hdr_purple = fill('4A235A')
fill_total = fill('1F4E79'); fill_gold = fill('D4AC0D')
fill_conc  = fill('D6E4F0'); fill_forma = fill('D5E8D4'); fill_aco = fill('EDE7F6')
seg_fills  = [fill('2E4057'), fill('1B4332'), fill('4A235A')]  # SEG1/2/3 subtotal
sub_fills  = [fill('EBF5FB'), fill('EAFAF1'), fill('F3EFF8')]  # SEG1/2/3 linhas

wf = Font(name='Arial', bold=True, color='FFFFFF', size=10)
nf = Font(name='Arial', size=10)
center  = Alignment(horizontal='center', vertical='center', wrap_text=True)
left_al = Alignment(horizontal='left',   vertical='center')
right_al= Alignment(horizontal='right',  vertical='center')
thin    = Side(style='thin', color='BFBFBF')
border  = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── PARSE QUANTITATIVOS ───────────────────────────────────────────────────────
def parse_qty(val):
    if pd.isna(val): return 0.0
    m = re.search(r'[\d,.]+', str(val).replace(',', '.'))
    return float(m.group()) if m else 0.0

def extract_bitola(desc):
    m = re.search(r'[Øø⌀-]?\s*([\d.]+)\s*mm', str(desc), re.IGNORECASE)
    return float(m.group(1)) if m else None

def load_qto(path):
    df = pd.read_excel(path, sheet_name=0)
    col_name  = [c for c in df.columns if 'Name'  in c][0]
    col_layer = [c for c in df.columns if 'Layer' in c and 'Layer Id' not in c and 'Name' not in c][0]
    desc_cols = sorted([c for c in df.columns if 'Descri' in c])
    qty_cols  = sorted([c for c in df.columns if 'Quantidade' in c])
    records = []
    for _, row in df.iterrows():
        rec = {'PILAR': str(row[col_name]).strip(),
               'CAMADA': str(row[col_layer]).strip(),
               'CONCRETO_m3': 0.0, 'FORMA_m2': 0.0}
        for dc, qc in zip(desc_cols, qty_cols):
            desc = str(row[dc]) if not pd.isna(row[dc]) else ''
            qty  = parse_qty(row[qc])
            if not desc or qty == 0: continue
            if 'Concreto' in desc:
                rec['CONCRETO_m3'] += qty
            elif 'Forma' in desc:
                rec['FORMA_m2'] += qty
            elif any(x in desc for x in ['Armadura','CA50','CA60']):
                b = extract_bitola(desc)
                if b:
                    k = f'ACO_{b}mm_kg'
                    rec[k] = rec.get(k, 0.0) + qty
        records.append(rec)
    result = pd.DataFrame(records)
    aco_cols = sorted({c for c in result.columns if c.startswith('ACO_')},
                      key=lambda x: float(re.search(r'ACO_([\d.]+)mm', x).group(1)))
    bitolas = [float(re.search(r'ACO_([\d.]+)mm', c).group(1))
               for c in aco_cols if result[c].sum() > 0]
    for c in [f'ACO_{b}mm_kg' for b in bitolas]:
        result[c] = result.get(c, 0.0)
    result[[f'ACO_{b}mm_kg' for b in bitolas]] = result[[f'ACO_{b}mm_kg' for b in bitolas]].fillna(0.0)
    return result, bitolas

def load_segmentos(seg_files):
    frames = []
    for f in seg_files:
        df = pd.read_excel(f, sheet_name=0)
        col_name = df.columns[0]
        col_seg  = df.columns[1]
        tmp = df[[col_name, col_seg]].copy()
        tmp.columns = ['PILAR', 'SEGMENTO']
        tmp['PILAR']    = tmp['PILAR'].astype(str).str.strip()
        tmp['SEGMENTO'] = tmp['SEGMENTO'].astype(str).str.strip()
        frames.append(tmp)
    return pd.concat(frames, ignore_index=True)

# ── HELPERS ───────────────────────────────────────────────────────────────────
def sc(ws, r, c, val, font=None, fill_=None, align=None, fmt=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font   = font  or nf
    cell.fill   = fill_ or PatternFill()
    cell.alignment = align or right_al
    cell.border = border
    if fmt: cell.number_format = fmt

def sh(ws, r, c, val, font=None, fill_=None):
    sc(ws, r, c, val, font=font or wf, fill_=fill_ or hdr_blue, align=center)

# ── ABA 1: DETALHADO (agrupado por segmento) ─────────────────────────────────
def build_detalhado(wb, result, bitolas):
    ws = wb.active; ws.title = 'DETALHADO'
    cols = ['PILAR','SEGMENTO','CAMADA','CONCRETO','FORMA'] + [f'ACO O{b}mm' for b in bitolas]
    units= ['','','','m3','m2'] + ['kg']*len(bitolas)
    hfills=[hdr_blue]*3+[hdr_blue,hdr_green]+[hdr_purple]*len(bitolas)
    for ci,(h,u,f) in enumerate(zip(cols,units,hfills),1):
        sh(ws,1,ci,h,fill_=f); sh(ws,2,ci,u,fill_=f)

    ri = 3
    for si, seg in enumerate(SEGMENTOS):
        sub = result[result['SEGMENTO']==seg]
        alt = [fill('F0F4F8'), fill('FFFFFF')]
        for idx,(_, row) in enumerate(sub.iterrows()):
            rf = alt[idx%2]
            vals = [row['PILAR'], seg, row['CAMADA'], row['CONCRETO_m3'], row['FORMA_m2']] + \
                   [row[f'ACO_{b}mm_kg'] for b in bitolas]
            for ci,v in enumerate(vals,1):
                fmt = '#,##0.00' if ci>3 else None
                sc(ws,ri,ci,v,fill_=rf, align=left_al if ci<=3 else right_al, fmt=fmt)
            ri += 1
        # subtotal segmento
        sf = seg_fills[si]
        stot = ['SUBTOTAL',seg,'',sub['CONCRETO_m3'].sum(),sub['FORMA_m2'].sum()] + \
               [sub[f'ACO_{b}mm_kg'].sum() for b in bitolas]
        for ci,v in enumerate(stot,1):
            c = ws.cell(row=ri,column=ci,value=v)
            c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
            c.fill=sf; c.border=border
            c.alignment=left_al if ci<=3 else right_al
            if ci>3 and isinstance(v,float): c.number_format='#,##0.00'
        ri += 1

    # total geral
    for ci,v in enumerate(['TOTAL GERAL','','',
                            result['CONCRETO_m3'].sum(),result['FORMA_m2'].sum()]+
                           [result[f'ACO_{b}mm_kg'].sum() for b in bitolas],1):
        c=ws.cell(row=ri,column=ci,value=v)
        c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
        c.fill=fill_total; c.border=border
        c.alignment=left_al if ci<=3 else right_al
        if ci>3 and isinstance(v,float): c.number_format='#,##0.00'

    for ci,w in enumerate([14,10,20,14,14]+[13]*len(bitolas),1):
        ws.column_dimensions[get_column_letter(ci)].width=w
    ws.row_dimensions[1].height=30; ws.row_dimensions[2].height=18
    ws.freeze_panes='A3'

# ── ABA 2: RESUMO TOTAL ───────────────────────────────────────────────────────
def build_resumo_total(wb, result, bitolas):
    ws2=wb.create_sheet('RESUMO TOTAL')
    ws2['A1']=f'RESUMO GERAL - {TITULO}'
    ws2['A1'].font=Font(name='Arial',bold=True,size=14,color='1F4E79')
    ws2.merge_cells('A1:D1'); ws2['A1'].alignment=center
    for ci,h in enumerate(['CATEGORIA','DESCRICAO','TOTAL','UNIDADE'],1):
        sh(ws2,3,ci,h)
    rows=[('CONCRETO','Concreto C-30 - Abatimento 12 cm',result['CONCRETO_m3'].sum(),'m3'),
          ('FORMA','Forma - Estrutura - Concreto',result['FORMA_m2'].sum(),'m2')]
    for b in bitolas:
        t=result[f'ACO_{b}mm_kg'].sum()
        if t>0: rows.append(('ACO',f'Armadura CA-Ø {b} mm',t,'kg'))
    fc={'CONCRETO':fill_conc,'FORMA':fill_forma,'ACO':fill_aco}
    for ri,(cat,desc,val,unit) in enumerate(rows,4):
        f=fc.get(cat,fill_conc)
        for ci,v in enumerate([cat,desc,val,unit],1):
            sc(ws2,ri,ci,v,fill_=f,align=left_al if ci==2 else center,
               fmt='#,##0.00' if ci==3 else None)
    r=4+len(rows)
    ag=sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
    for ci,v in enumerate(['ACO TOTAL','Total Geral ACO CA50+CA60',ag,'kg'],1):
        c=ws2.cell(row=r,column=ci,value=v)
        c.font=Font(name='Arial',bold=True,size=10); c.fill=fill_gold
        c.border=border; c.alignment=left_al if ci==2 else center
        if ci==3: c.number_format='#,##0.00'
    ws2.column_dimensions['A'].width=14; ws2.column_dimensions['B'].width=44
    ws2.column_dimensions['C'].width=16; ws2.column_dimensions['D'].width=12

# ── ABA 3: RESUMO POR SEGMENTO ────────────────────────────────────────────────
def build_resumo_segmento(wb, result, bitolas):
    ws3=wb.create_sheet('RESUMO POR SEGMENTO')
    nc=2+len(bitolas)+2
    ws3['A1']=f'RESUMO POR SEGMENTO - {TITULO}'
    ws3['A1'].font=Font(name='Arial',bold=True,size=14,color='1F4E79')
    ws3.merge_cells(f'A1:{get_column_letter(nc)}1'); ws3['A1'].alignment=center
    hdrs=['SEGMENTO','CONCRETO (m3)','FORMA (m2)']+[f'ACO O{b}mm (kg)' for b in bitolas]+['ACO TOTAL (kg)']
    for ci,h in enumerate(hdrs,1): sh(ws3,3,ci,h)

    for ri,seg in enumerate(SEGMENTOS,4):
        sub=result[result['SEGMENTO']==seg]
        at=sum(sub[f'ACO_{b}mm_kg'].sum() for b in bitolas)
        sf=seg_fills[ri-4]
        row=[seg,sub['CONCRETO_m3'].sum(),sub['FORMA_m2'].sum()]+\
            [sub[f'ACO_{b}mm_kg'].sum() for b in bitolas]+[at]
        for ci,v in enumerate(row,1):
            c=ws3.cell(row=ri,column=ci,value=v)
            c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
            c.fill=sf; c.border=border
            c.alignment=left_al if ci==1 else right_al
            if ci>1: c.number_format='#,##0.00'

    rtt=4+len(SEGMENTOS)
    aa=sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
    tt=['TOTAL GERAL',result['CONCRETO_m3'].sum(),result['FORMA_m2'].sum()]+\
       [result[f'ACO_{b}mm_kg'].sum() for b in bitolas]+[aa]
    for ci,v in enumerate(tt,1):
        c=ws3.cell(row=rtt,column=ci,value=v)
        c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
        c.fill=fill_total; c.border=border
        c.alignment=left_al if ci==1 else right_al
        if ci>1: c.number_format='#,##0.00'

    for ci,w in enumerate([16,16,16]+[14]*len(bitolas)+[16],1):
        ws3.column_dimensions[get_column_letter(ci)].width=w

# ── ABA 4: PRORRATA ───────────────────────────────────────────────────────────
def build_prorrata(wb, result, bitolas):
    ws4=wb.create_sheet('PRORRATA')
    tot_conc=result['CONCRETO_m3'].sum()
    tot_forma=result['FORMA_m2'].sum()
    tot_acos={b:result[f'ACO_{b}mm_kg'].sum() for b in bitolas}
    tot_aco_g=sum(tot_acos.values())

    col_map=[]
    col_map.append(('PILAR',   '',  hdr_blue,  False,'PILAR',        None))
    col_map.append(('SEGMENTO','',  hdr_blue,  False,'SEGMENTO',     None))
    col_map.append(('CONCRETO','m3',hdr_blue,  False,'CONCRETO_m3',  tot_conc))
    col_map.append(('%',       '%', hdr_blue,  True, 'CONCRETO_m3',  tot_conc))
    col_map.append(('FORMA',   'm2',hdr_green, False,'FORMA_m2',     tot_forma))
    col_map.append(('%',       '%', hdr_green, True, 'FORMA_m2',     tot_forma))
    for b in bitolas:
        col_map.append((f'ACO O{b}mm','kg',hdr_purple,False,f'ACO_{b}mm_kg',tot_acos[b]))
        col_map.append(('%',          '%', hdr_purple,True, f'ACO_{b}mm_kg',tot_acos[b]))
    col_map.append(('ACO TOTAL','kg',fill_gold,False,'__ACO_TOTAL__',tot_aco_g))
    col_map.append(('%',        '%', fill_gold,True, '__ACO_TOTAL__',tot_aco_g))

    nc=len(col_map)
    ws4['A1']=f'PRORRATA (% PARTICIPAÇÃO) - {TITULO}'
    ws4['A1'].font=Font(name='Arial',bold=True,size=14,color='1F4E79')
    ws4.merge_cells(f'A1:{get_column_letter(nc)}1'); ws4['A1'].alignment=center
    for ci,(l2,l3,f,_,__,___) in enumerate(col_map,1):
        sh(ws4,2,ci,l2,fill_=f); sh(ws4,3,ci,l3,fill_=f)

    ri=4
    for si,seg in enumerate(SEGMENTOS):
        sub=result[result['SEGMENTO']==seg]
        alt=[sub_fills[si], fill('FFFFFF')]
        for idx,(_,row) in enumerate(sub.iterrows()):
            rf=alt[idx%2]
            aco_row=sum(row[f'ACO_{b}mm_kg'] for b in bitolas)
            for ci,(_,_,_,is_pct,key,tot) in enumerate(col_map,1):
                if key=='PILAR':
                    sc(ws4,ri,ci,row['PILAR'],fill_=rf,align=left_al)
                elif key=='SEGMENTO':
                    sc(ws4,ri,ci,seg,fill_=rf,align=left_al)
                elif key=='__ACO_TOTAL__':
                    val=(aco_row/tot*100) if (is_pct and tot) else aco_row
                    sc(ws4,ri,ci,val,fill_=rf,fmt='0.00"%"' if is_pct else '#,##0.00')
                else:
                    av=row.get(key,0.0)
                    val=(av/tot*100) if (is_pct and tot) else av
                    sc(ws4,ri,ci,val,fill_=rf,fmt='0.00"%"' if is_pct else '#,##0.00')
            ri+=1

    # total geral
    aco_g=sum(result[f'ACO_{b}mm_kg'].sum() for b in bitolas)
    for ci,(_,_,_,is_pct,key,tot) in enumerate(col_map,1):
        c=ws4.cell(row=ri,column=ci)
        c.border=border
        c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
        c.fill=fill_total
        if key=='PILAR':
            c.value='TOTAL'; c.alignment=left_al
        elif key=='SEGMENTO':
            c.value=''; c.alignment=left_al
        elif is_pct:
            c.value=100.0; c.number_format='0.00"%"'; c.alignment=right_al
        elif key=='__ACO_TOTAL__':
            c.value=aco_g; c.number_format='#,##0.00'; c.alignment=right_al
        else:
            c.value=tot; c.number_format='#,##0.00'; c.alignment=right_al

    for ci,(_,_,_,is_pct,__,___) in enumerate(col_map,1):
        if ci<=2: ws4.column_dimensions[get_column_letter(ci)].width=14
        elif is_pct: ws4.column_dimensions[get_column_letter(ci)].width=8
        else: ws4.column_dimensions[get_column_letter(ci)].width=13
    ws4.row_dimensions[2].height=30; ws4.row_dimensions[3].height=18
    ws4.freeze_panes='A4'

# ── MAIN ──────────────────────────────────────────────────────────────────────
def main():
    print(f"Carregando quantitativos: {FILE_QTO.name}")
    result, bitolas = load_qto(FILE_QTO)

    print("Carregando segmentos...")
    seg_map = load_segmentos(SEG_FILES)

    # merge — join por PILAR
    result = result.merge(seg_map[['PILAR','SEGMENTO']], on='PILAR', how='left')
    sem_seg = result['SEGMENTO'].isna().sum()
    if sem_seg:
        print(f"  ⚠ {sem_seg} pilares sem segmento definido → classificados como 'SEM SEG'")
        result['SEGMENTO'] = result['SEGMENTO'].fillna('SEM SEG')

    # ordenar por segmento na ordem definida
    seg_order = {s:i for i,s in enumerate(SEGMENTOS)}
    result['_ord'] = result['SEGMENTO'].map(seg_order).fillna(99)
    result = result.sort_values(['_ord','PILAR']).drop(columns='_ord').reset_index(drop=True)

    print(f"  {len(result)} pilares | bitolas: {bitolas}")
    for seg in SEGMENTOS:
        n=len(result[result['SEGMENTO']==seg])
        print(f"  {seg}: {n} pilares")

    wb=Workbook()
    build_detalhado(wb, result, bitolas)
    build_resumo_total(wb, result, bitolas)
    build_resumo_segmento(wb, result, bitolas)
    build_prorrata(wb, result, bitolas)
    wb.save(FILE_OUT)
    print(f"\nSalvo: {FILE_OUT}")

if __name__=='__main__':
    main()
