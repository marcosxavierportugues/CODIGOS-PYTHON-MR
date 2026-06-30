# -*- coding: utf-8 -*-
"""
Conferidor de listas de materiais Navisworks x planilha tratada.

Uso por terminal:
    python conferir_lista_materiais.py "navis.xlsx" "tratada.xlsx"

Se rodar sem argumentos, abre janelas para escolher:
1) planilha bruta do Navisworks/AltoQi
2) planilha ja tratada pelo gerador
3) local do relatorio de conferencia

O relatorio gerado possui 13 testes e abas com divergencias e dados normalizados.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ZERO = Decimal("0")
NUMERIC_EPS = Decimal("0.000001")
BASE_MATERIAL_ORDER = ["CONCRETO", "FORMA"]
STATUS_ORDER = {"ERRO": 0, "ALERTA": 1, "OK": 2}


@dataclass
class ParsedQuantity:
    value: Decimal
    ok: bool
    token: str = ""
    warning: str = ""
    error: str = ""


@dataclass
class SheetDetection:
    kind: str
    score: int
    details: str


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip().casefold()


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")
    return str(value)


def decimal_close(left: Decimal, right: Decimal, eps: Decimal = NUMERIC_EPS) -> bool:
    return abs(left - right) <= eps


def decimal_to_report(value: Decimal | None) -> Any:
    if value is None:
        return ""
    if value == value.to_integral():
        return int(value)
    return float(value)


def parse_quantity(value: Any) -> ParsedQuantity:
    """Le quantidade com decimal brasileiro/internacional e ignora unidades."""
    if value is None:
        return ParsedQuantity(ZERO, True)
    if isinstance(value, bool):
        return ParsedQuantity(Decimal(int(value)), True, token=str(value))
    if isinstance(value, int):
        return ParsedQuantity(Decimal(value), True, token=str(value))
    if isinstance(value, float):
        return ParsedQuantity(Decimal(str(value)), True, token=str(value))

    text = str(value).strip()
    if not text:
        return ParsedQuantity(ZERO, True)

    match = re.search(r"[+-]?\d(?:[\d\s.,]*\d)?", text)
    if not match:
        return ParsedQuantity(
            ZERO,
            False,
            error=f"Quantidade sem numero reconhecivel: {text!r}",
        )

    token = re.sub(r"\s+", "", match.group(0))
    original_token = token
    sign = ""
    if token.startswith(("+", "-")):
        sign, token = token[0], token[1:]

    warning_parts = []
    if re.search(r"[A-Za-z]", text):
        warning_parts.append("quantidade com unidade/texto")
    if "," in token:
        warning_parts.append("usa virgula decimal ou milhar")
    if token.count(".") + token.count(",") > 1:
        warning_parts.append("possui multiplos separadores")

    if "," in token and "." in token:
        decimal_separator = "," if token.rfind(",") > token.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        token = token.replace(thousands_separator, "")
        token = token.replace(decimal_separator, ".")
    elif "," in token:
        parts = token.split(",")
        if len(parts) > 2 and all(len(part) == 3 for part in parts[1:]):
            token = "".join(parts)
        else:
            token = "".join(parts[:-1]) + "." + parts[-1]
    elif token.count(".") > 1:
        parts = token.split(".")
        if all(len(part) == 3 for part in parts[1:]):
            token = "".join(parts)
        else:
            token = "".join(parts[:-1]) + "." + parts[-1]

    try:
        return ParsedQuantity(
            Decimal(sign + token),
            True,
            token=original_token,
            warning="; ".join(warning_parts),
        )
    except InvalidOperation:
        return ParsedQuantity(
            ZERO,
            False,
            token=original_token,
            error=f"Quantidade invalida: {text!r}",
        )


def bitola_label(value: Decimal) -> str:
    value = value.quantize(Decimal("0.1"))
    return f"O{value:.1f}mm"


def extract_bitola(description: Any) -> Decimal | None:
    text = str(description or "").replace(",", ".")
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*mm", text, flags=re.IGNORECASE)
    if not match:
        return None
    try:
        return Decimal(match.group(1)).quantize(Decimal("0.1"))
    except InvalidOperation:
        return None


def classify_description(description: Any) -> str | None:
    normalized = normalize_text(description)
    if not normalized:
        return None
    if "forma" in normalized:
        return "FORMA"
    if "concreto" in normalized:
        return "CONCRETO"
    if any(term in normalized for term in ("armadura", "aco", "ca50", "ca60")):
        bitola = extract_bitola(description)
        return bitola_label(bitola) if bitola is not None else "ACO_SEM_BITOLA"
    return "NAO_RECONHECIDO"


def material_sort_key(material: str) -> tuple[int, Decimal, str]:
    if material == "CONCRETO":
        return (0, ZERO, material)
    if material == "FORMA":
        return (1, ZERO, material)
    match = re.search(r"O([0-9]+(?:\.[0-9]+)?)mm", material)
    if match:
        return (2, Decimal(match.group(1)), material)
    return (3, ZERO, material)


def detect_workbook(path: Path) -> SheetDetection:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # pragma: no cover - defensive for corrupted files
        return SheetDetection("ERRO", -100, str(exc))

    score_raw = 0
    score_treated = 0
    details = []

    sheet_names_norm = {normalize_text(name): name for name in workbook.sheetnames}
    if any("todos" == key or key.startswith("todos") for key in sheet_names_norm):
        score_treated += 3
    if any("total geral" in key or key.startswith("total") for key in sheet_names_norm):
        score_treated += 3

    for worksheet in workbook.worksheets:
        headers = [
            normalize_text(worksheet.cell(1, col).value)
            for col in range(1, min(worksheet.max_column, 30) + 1)
        ]
        joined = " | ".join(headers)
        if "descricao" in joined or "descrição" in joined:
            score_raw += 2
        if "quantidade" in joined or "quantity" in joined:
            score_raw += 2
        if "elemento" in joined and "concreto" in joined and "forma" in joined:
            score_treated += 4
        if "item" in joined and "name" in joined:
            score_raw += 1

    if score_treated > score_raw:
        kind = "TRATADA"
        score = score_treated
    elif score_raw > score_treated:
        kind = "BRUTA"
        score = score_raw
    else:
        kind = "INDEFINIDA"
        score = score_raw
    details.append(f"score_bruta={score_raw}; score_tratada={score_treated}")
    return SheetDetection(kind, score, "; ".join(details))


def find_raw_columns(worksheet: Any) -> tuple[int, list[tuple[int, int]], list[str]]:
    headers = [worksheet.cell(1, col).value for col in range(1, worksheet.max_column + 1)]
    normalized = [normalize_text(header) for header in headers]

    item_col = 1
    for idx, header in enumerate(normalized, start=1):
        if "item" in header[:20] or header.endswith("name") or " name" in header:
            item_col = idx
            break

    desc_cols = []
    qty_cols = []
    for idx, header in enumerate(normalized, start=1):
        if "descricao" in header or "description" in header:
            desc_cols.append(idx)
        if "quantidade" in header or "quantity" in header:
            qty_cols.append(idx)

    pairs: list[tuple[int, int]] = []
    used_qty = set()
    for desc_col in desc_cols:
        chosen = None
        if desc_col + 1 in qty_cols:
            chosen = desc_col + 1
        else:
            for qty_col in qty_cols:
                if qty_col not in used_qty and qty_col > desc_col:
                    chosen = qty_col
                    break
        if chosen is not None:
            pairs.append((desc_col, chosen))
            used_qty.add(chosen)

    messages = []
    if not pairs:
        messages.append("Nenhum par Descricao/Quantidade encontrado.")
    return item_col, pairs, messages


def read_raw(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    worksheet = workbook[workbook.sheetnames[0]]
    item_col, pairs, setup_messages = find_raw_columns(worksheet)

    rows = []
    issues = []
    materials = set()

    for row_index in range(2, worksheet.max_row + 1):
        item_value = worksheet.cell(row_index, item_col).value
        item = str(item_value).strip() if item_value is not None else ""
        if not item:
            continue

        record = {
            "row": row_index,
            "elemento": item,
            "nivel": "",
            "values": {},
        }

        for desc_col, qty_col in pairs:
            description = worksheet.cell(row_index, desc_col).value
            quantity_raw = worksheet.cell(row_index, qty_col).value
            if description in (None, "") and quantity_raw in (None, ""):
                continue
            if description in (None, "") and quantity_raw not in (None, ""):
                issues.append(
                    issue(
                        "T09",
                        "ERRO",
                        path,
                        worksheet.title,
                        row_index,
                        qty_col,
                        item,
                        "",
                        "",
                        quantity_raw,
                        "Quantidade preenchida sem descricao.",
                    )
                )
                continue
            if description not in (None, "") and quantity_raw in (None, ""):
                issues.append(
                    issue(
                        "T09",
                        "ERRO",
                        path,
                        worksheet.title,
                        row_index,
                        qty_col,
                        item,
                        "",
                        "",
                        "",
                        "Descricao preenchida sem quantidade.",
                    )
                )
                continue

            material = classify_description(description)
            if material in (None, "NAO_RECONHECIDO", "ACO_SEM_BITOLA"):
                issues.append(
                    issue(
                        "T08",
                        "ERRO",
                        path,
                        worksheet.title,
                        row_index,
                        desc_col,
                        item,
                        material or "",
                        "Concreto, Forma ou Armadura com bitola",
                        description,
                        "Descricao nao reconhecida.",
                    )
                )
                continue

            parsed = parse_quantity(quantity_raw)
            if not parsed.ok:
                issues.append(
                    issue(
                        "T09",
                        "ERRO",
                        path,
                        worksheet.title,
                        row_index,
                        qty_col,
                        item,
                        material,
                        "Numero",
                        quantity_raw,
                        parsed.error,
                    )
                )
                continue
            if parsed.warning:
                issues.append(
                    issue(
                        "T09",
                        "ALERTA",
                        path,
                        worksheet.title,
                        row_index,
                        qty_col,
                        item,
                        material,
                        "",
                        quantity_raw,
                        parsed.warning,
                    )
                )

            record["values"][material] = record["values"].get(material, ZERO) + parsed.value
            materials.add(material)

        rows.append(record)

    for message in setup_messages:
        issues.append(
            issue("T03", "ERRO", path, worksheet.title, 1, 1, "", "", "", "", message)
        )

    return {
        "path": path,
        "sheet": worksheet.title,
        "rows": rows,
        "materials": materials,
        "issues": issues,
        "desc_qty_pairs": pairs,
        "item_col": item_col,
    }


def find_todos_sheet(workbook: Any) -> Any:
    for name in workbook.sheetnames:
        if normalize_text(name) == "todos" or normalize_text(name).startswith("todos"):
            return workbook[name]
    for worksheet in workbook.worksheets:
        for row_index in range(1, min(10, worksheet.max_row) + 1):
            headers = [normalize_text(worksheet.cell(row_index, col).value) for col in range(1, worksheet.max_column + 1)]
            joined = " | ".join(headers)
            if "elemento" in joined and "concreto" in joined and "forma" in joined:
                return worksheet
    raise ValueError("Aba TODOS nao encontrada.")


def find_header_row(worksheet: Any, required_terms: list[str]) -> int:
    for row_index in range(1, min(20, worksheet.max_row) + 1):
        values = [normalize_text(worksheet.cell(row_index, col).value) for col in range(1, worksheet.max_column + 1)]
        joined = " | ".join(values)
        if all(term in joined for term in required_terms):
            return row_index
    raise ValueError(f"Cabecalho nao encontrado na aba {worksheet.title}.")


def treated_header_key(header: Any) -> str | None:
    text = str(header or "").strip()
    normalized = normalize_text(text)
    if not normalized:
        return None
    if "elemento" in normalized or normalized == "item":
        return "ELEMENTO"
    if "nivel" in normalized or "pav" in normalized:
        return "NIVEL"
    if "concreto" in normalized:
        return "CONCRETO"
    if "forma" in normalized:
        return "FORMA"
    match = re.search(r"[O0Øø]\s*([0-9]+(?:[,.][0-9]+)?)\s*mm", text, flags=re.IGNORECASE)
    if match:
        return bitola_label(Decimal(match.group(1).replace(",", ".")))
    return None


def read_treated(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    todos = find_todos_sheet(workbook)
    header_row = find_header_row(todos, ["concreto", "forma"])

    column_map: dict[str, int] = {}
    issues = []
    materials = set()
    for col in range(1, todos.max_column + 1):
        key = treated_header_key(todos.cell(header_row, col).value)
        if key:
            column_map[key] = col
            if key not in ("ELEMENTO", "NIVEL"):
                materials.add(key)

    for required in ("ELEMENTO", "CONCRETO", "FORMA"):
        if required not in column_map:
            issues.append(
                issue("T04", "ERRO", path, todos.title, header_row, 1, "", required, required, "", "Coluna obrigatoria nao encontrada.")
            )

    data_rows = []
    total_rows = []
    data_start = header_row + 1
    for row_index in range(data_start, todos.max_row + 1):
        first_value = todos.cell(row_index, column_map.get("ELEMENTO", 1)).value
        first_text = str(first_value).strip() if first_value is not None else ""
        first_norm = normalize_text(first_text)
        row_values = [todos.cell(row_index, col).value for col in range(1, todos.max_column + 1)]
        if all(value in (None, "") for value in row_values):
            continue
        if not first_text and any(normalize_text(v) in {"m3", "m2", "kg"} for v in row_values):
            continue

        is_total = first_norm.startswith("total") or first_norm.startswith("subtotal")
        record = {
            "row": row_index,
            "elemento": first_text,
            "nivel": str(todos.cell(row_index, column_map.get("NIVEL", 0)).value or "").strip() if "NIVEL" in column_map else "",
            "values": {},
        }

        for material in sorted(materials, key=material_sort_key):
            col = column_map[material]
            raw_value = todos.cell(row_index, col).value
            if raw_value in (None, ""):
                parsed = ParsedQuantity(ZERO, True)
            else:
                parsed = parse_quantity(raw_value)
            if not parsed.ok:
                issues.append(
                    issue(
                        "T10",
                        "ERRO",
                        path,
                        todos.title,
                        row_index,
                        col,
                        first_text,
                        material,
                        "Numero",
                        raw_value,
                        parsed.error,
                    )
                )
                parsed = ParsedQuantity(ZERO, True)
            elif parsed.value < ZERO:
                issues.append(
                    issue(
                        "T10",
                        "ERRO",
                        path,
                        todos.title,
                        row_index,
                        col,
                        first_text,
                        material,
                        ">= 0",
                        raw_value,
                        "Valor negativo.",
                    )
                )
            record["values"][material] = parsed.value

        if is_total:
            total_rows.append(record)
        elif first_text:
            data_rows.append(record)

    summary = read_total_geral(workbook, path, materials)
    return {
        "path": path,
        "sheet": todos.title,
        "rows": data_rows,
        "total_rows": total_rows,
        "materials": materials,
        "issues": issues + summary["issues"],
        "column_map": column_map,
        "summary": summary,
    }


def read_total_geral(workbook: Any, path: Path, known_materials: set[str]) -> dict[str, Any]:
    worksheet = None
    for name in workbook.sheetnames:
        normalized = normalize_text(name)
        if "total geral" in normalized or normalized.startswith("total"):
            worksheet = workbook[name]
            break
    if worksheet is None:
        return {"sheet": "", "totals": {}, "units": {}, "total_aco": None, "issues": [issue("T04", "ERRO", path, "", 0, 0, "", "", "Aba TOTAL GERAL", "", "Aba TOTAL GERAL nao encontrada.")]}

    issues = []
    totals: dict[str, Decimal] = {}
    units: dict[str, str] = {}
    total_aco = None

    try:
        header_row = find_header_row(worksheet, ["total geral", "descricao"])
    except ValueError as exc:
        return {"sheet": worksheet.title, "totals": totals, "units": units, "total_aco": None, "issues": [issue("T04", "ERRO", path, worksheet.title, 0, 0, "", "", "", "", str(exc))]}

    header_lookup = {}
    for col in range(1, worksheet.max_column + 1):
        key = normalize_text(worksheet.cell(header_row, col).value)
        if key:
            header_lookup[key] = col

    desc_col = next((col for key, col in header_lookup.items() if "descricao" in key), 3)
    total_col = next((col for key, col in header_lookup.items() if key == "total geral" or key.startswith("total geral")), 4)
    unit_col = next((col for key, col in header_lookup.items() if key in {"un", "unidade"}), None)

    for row_index in range(header_row + 1, worksheet.max_row + 1):
        row_text = " ".join(str(worksheet.cell(row_index, col).value or "") for col in range(1, worksheet.max_column + 1))
        if not row_text.strip():
            continue
        first_text = normalize_text(worksheet.cell(row_index, 1).value)
        description = worksheet.cell(row_index, desc_col).value
        material = classify_summary_description(description, first_text, known_materials)
        if first_text.startswith("total aco"):
            parsed = parse_quantity(worksheet.cell(row_index, total_col).value)
            if parsed.ok:
                total_aco = parsed.value
            continue
        if not material:
            continue
        parsed = parse_quantity(worksheet.cell(row_index, total_col).value)
        if not parsed.ok:
            issues.append(
                issue("T13", "ERRO", path, worksheet.title, row_index, total_col, "", material, "Numero", worksheet.cell(row_index, total_col).value, parsed.error)
            )
            continue
        totals[material] = parsed.value
        if unit_col:
            units[material] = str(worksheet.cell(row_index, unit_col).value or "").strip()

    return {
        "sheet": worksheet.title,
        "totals": totals,
        "units": units,
        "total_aco": total_aco,
        "issues": issues,
    }


def classify_summary_description(description: Any, first_text: str, known_materials: set[str]) -> str | None:
    material = classify_description(description)
    if material and material not in {"NAO_RECONHECIDO", "ACO_SEM_BITOLA"}:
        return material
    if "concreto" in first_text:
        return "CONCRETO"
    if "forma" in first_text:
        return "FORMA"
    if "aco" in first_text:
        text = str(description or "")
        for known in known_materials:
            if known.startswith("O") and known[1:-2] in text:
                return known
    return None


def issue(
    test_id: str,
    severity: str,
    path: Path | str,
    sheet: str,
    row: int,
    column: int,
    elemento: str,
    material: str,
    expected: Any,
    found: Any,
    detail: str,
) -> dict[str, Any]:
    return {
        "teste": test_id,
        "severidade": severity,
        "arquivo": str(path),
        "aba": sheet,
        "linha": row or "",
        "coluna": column or "",
        "elemento": elemento,
        "material": material,
        "esperado": display_value(expected),
        "encontrado": display_value(found),
        "detalhe": detail,
    }


def sum_material(rows: list[dict[str, Any]], material: str) -> Decimal:
    return sum((row["values"].get(material, ZERO) for row in rows), ZERO)


def expected_legacy_value(material: str, raw_value: Decimal) -> Decimal:
    """Regra vista em planilhas antigas com separador decimal perdido."""
    if material in {"CONCRETO", "FORMA"}:
        return (raw_value * Decimal("1000")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    text = format(raw_value.normalize(), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    sign = "-" if text.startswith("-") else ""
    digits = text.lstrip("-").replace(".", "").replace(",", "")
    return Decimal(sign + (digits or "0"))


def compare_rows(raw: dict[str, Any], treated: dict[str, Any], materials: list[str]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int, int]:
    real_issues = []
    legacy_issues = []
    real_matches = 0
    legacy_matches = 0

    row_count = min(len(raw["rows"]), len(treated["rows"]))
    for index in range(row_count):
        raw_row = raw["rows"][index]
        treated_row = treated["rows"][index]
        for material in materials:
            raw_value = raw_row["values"].get(material, ZERO)
            treated_value = treated_row["values"].get(material, ZERO)
            if decimal_close(raw_value, treated_value):
                real_matches += 1
            else:
                real_issues.append(
                    issue(
                        "T11",
                        "ERRO",
                        treated["path"],
                        treated["sheet"],
                        treated_row["row"],
                        "",
                        treated_row["elemento"],
                        material,
                        raw_value,
                        treated_value,
                        f"Linha {index + 1}: valor tratado nao bate com valor real da bruta.",
                    )
                )

            legacy_value = expected_legacy_value(material, raw_value)
            if decimal_close(legacy_value, treated_value):
                legacy_matches += 1
            else:
                legacy_issues.append(
                    issue(
                        "T12",
                        "ERRO",
                        treated["path"],
                        treated["sheet"],
                        treated_row["row"],
                        "",
                        treated_row["elemento"],
                        material,
                        legacy_value,
                        treated_value,
                        f"Linha {index + 1}: valor tambem nao bate com a regra antiga/escalada.",
                    )
                )

    return real_issues, legacy_issues, real_matches, legacy_matches


def run_tests(raw_path: Path, treated_path: Path) -> dict[str, Any]:
    all_issues = []
    tests = []

    original_raw_path = raw_path
    original_treated_path = treated_path
    det1 = detect_workbook(raw_path)
    det2 = detect_workbook(treated_path)
    original_det1 = det1
    original_det2 = det2
    swapped = det1.kind == "TRATADA" and det2.kind == "BRUTA"
    if swapped:
        raw_path, treated_path = treated_path, raw_path
        det1, det2 = det2, det1

    add_test(
        tests,
        "T01",
        "Arquivos existem e abrem",
        "OK" if raw_path.exists() and treated_path.exists() else "ERRO",
        f"Bruta: {raw_path}; Tratada: {treated_path}",
    )
    add_test(
        tests,
        "T02",
        "Ordem dos arquivos / deteccao automatica",
        "ALERTA" if swapped else ("OK" if det1.kind == "BRUTA" and det2.kind == "TRATADA" else "ERRO"),
        (
            f"Entrada 1={original_det1.kind} ({original_raw_path}); "
            f"Entrada 2={original_det2.kind} ({original_treated_path}); "
            f"usado como bruta={raw_path}; usado como tratada={treated_path}."
        )
        + (" Os caminhos foram invertidos automaticamente." if swapped else ""),
    )

    raw = read_raw(raw_path)
    treated = read_treated(treated_path)
    all_issues.extend(raw["issues"])
    all_issues.extend(treated["issues"])

    materials = sorted(raw["materials"].union(treated["materials"]), key=material_sort_key)

    add_test(
        tests,
        "T03",
        "Colunas da planilha bruta",
        "OK" if raw["desc_qty_pairs"] and not has_error(raw["issues"], "T03") else "ERRO",
        f"Item col={raw['item_col']}; pares Descricao/Quantidade={len(raw['desc_qty_pairs'])}.",
    )
    add_test(
        tests,
        "T04",
        "Estrutura da planilha tratada",
        "OK" if not has_error(treated["issues"], "T04") else "ERRO",
        f"Aba TODOS={treated['sheet']}; Aba TOTAL GERAL={treated['summary']['sheet']}; materiais={', '.join(materials)}.",
    )
    add_test(
        tests,
        "T05",
        "Quantidade de linhas de itens",
        "OK" if len(raw["rows"]) == len(treated["rows"]) else "ERRO",
        f"Bruta={len(raw['rows'])}; Tratada={len(treated['rows'])}.",
    )

    raw_elements = [row["elemento"] for row in raw["rows"]]
    treated_elements = [row["elemento"] for row in treated["rows"]]
    order_ok = raw_elements == treated_elements
    if not order_ok:
        for index, (left, right) in enumerate(zip(raw_elements, treated_elements), start=1):
            if left != right:
                row_number = treated["rows"][index - 1]["row"] if index - 1 < len(treated["rows"]) else ""
                all_issues.append(
                    issue("T06", "ERRO", treated_path, treated["sheet"], row_number, "", right, "", left, right, f"Elemento fora de ordem na linha comparativa {index}.")
                )
                break
    add_test(
        tests,
        "T06",
        "Ordem dos elementos",
        "OK" if order_ok else "ERRO",
        "Compara item a item e preserva duplicados.",
    )

    duplicates_ok = Counter(raw_elements) == Counter(treated_elements)
    add_test(
        tests,
        "T07",
        "Elementos duplicados preservados",
        "OK" if duplicates_ok else "ERRO",
        f"Duplicados bruta={count_duplicates(raw_elements)}; duplicados tratada={count_duplicates(treated_elements)}.",
    )

    add_test(
        tests,
        "T08",
        "Descricoes e bitolas reconhecidas",
        "OK" if not has_error(raw["issues"], "T08") else "ERRO",
        "Confere Concreto, Forma e Armadura com bitola em mm.",
    )
    add_test(
        tests,
        "T09",
        "Quantidades da bruta parseaveis",
        test_status_from_issues(raw["issues"], "T09"),
        "Tambem avisa celulas com unidade/texto, virgula ou multiplos separadores.",
    )
    add_test(
        tests,
        "T10",
        "Numeros da tratada validos",
        test_status_from_issues(treated["issues"], "T10"),
        "Campos vazios contam como zero; valores negativos ou texto invalido viram erro.",
    )

    real_issues, legacy_issues, real_matches, legacy_matches = compare_rows(raw, treated, materials)
    all_issues.extend(real_issues)
    total_compares = min(len(raw["rows"]), len(treated["rows"])) * len(materials)
    real_ok = not real_issues and len(raw["rows"]) == len(treated["rows"])
    add_test(
        tests,
        "T11",
        "Valores item a item - valor real",
        "OK" if real_ok else "ERRO",
        f"Iguais={real_matches}/{total_compares}. Este e o teste principal para saber se a tratada manteve o valor numerico real.",
    )

    legacy_ok = not legacy_issues and len(raw["rows"]) == len(treated["rows"])
    if not real_ok and legacy_ok:
        legacy_status = "ALERTA"
        legacy_detail = "A tratada nao bate com valor real, mas bate com a regra antiga/escalada; provavel perda de separador decimal."
    elif legacy_ok:
        legacy_status = "OK"
        legacy_detail = "A regra antiga tambem bateu, mas o valor real ja estava correto."
    else:
        legacy_status = "ERRO"
        legacy_detail = f"Divergencias tambem na regra antiga/escalada: {len(legacy_issues)}."
        all_issues.extend(legacy_issues)
    add_test(
        tests,
        "T12",
        "Separador decimal / escala antiga",
        legacy_status,
        legacy_detail,
    )

    total_issues = check_totals(treated, materials)
    all_issues.extend(total_issues)
    add_test(
        tests,
        "T13",
        "Totais TODOS e TOTAL GERAL",
        "OK" if not total_issues else "ERRO",
        "Soma linhas da aba TODOS, compara com linha TOTAL, aba TOTAL GERAL, unidades e TOTAL ACO.",
    )

    return {
        "raw": raw,
        "treated": treated,
        "materials": materials,
        "tests": tests,
        "issues": all_issues,
        "swapped": swapped,
    }


def add_test(tests: list[dict[str, Any]], test_id: str, name: str, status: str, detail: str) -> None:
    tests.append({"teste": test_id, "nome": name, "status": status, "detalhe": detail})


def has_error(issues: list[dict[str, Any]], test_id: str) -> bool:
    return any(item["teste"] == test_id and item["severidade"] == "ERRO" for item in issues)


def test_status_from_issues(issues: list[dict[str, Any]], test_id: str) -> str:
    relevant = [item for item in issues if item["teste"] == test_id]
    if any(item["severidade"] == "ERRO" for item in relevant):
        return "ERRO"
    if any(item["severidade"] == "ALERTA" for item in relevant):
        return "ALERTA"
    return "OK"


def count_duplicates(items: list[str]) -> int:
    return sum(count - 1 for count in Counter(items).values() if count > 1)


def check_totals(treated: dict[str, Any], materials: list[str]) -> list[dict[str, Any]]:
    issues = []
    path = treated["path"]
    sheet = treated["sheet"]
    data_rows = treated["rows"]
    total_rows = treated["total_rows"]
    summary = treated["summary"]

    if not total_rows:
        issues.append(issue("T13", "ERRO", path, sheet, "", "", "", "", "Linha TOTAL", "", "Aba TODOS nao possui linha TOTAL."))
    total_row = total_rows[-1] if total_rows else None

    for material in materials:
        data_sum = sum_material(data_rows, material)
        if total_row is not None:
            total_value = total_row["values"].get(material, ZERO)
            if not decimal_close(data_sum, total_value):
                issues.append(
                    issue("T13", "ERRO", path, sheet, total_row["row"], "", "TOTAL", material, data_sum, total_value, "Linha TOTAL da aba TODOS nao bate com soma dos itens.")
                )

        summary_value = summary["totals"].get(material)
        if summary_value is None:
            issues.append(
                issue("T13", "ERRO", path, summary["sheet"], "", "", "", material, data_sum, "", "Material nao encontrado na aba TOTAL GERAL.")
            )
        elif not decimal_close(data_sum, summary_value):
            issues.append(
                issue("T13", "ERRO", path, summary["sheet"], "", "", "", material, data_sum, summary_value, "TOTAL GERAL nao bate com soma dos itens.")
            )

        expected_unit = "m3" if material == "CONCRETO" else "m2" if material == "FORMA" else "kg"
        found_unit = normalize_text(summary["units"].get(material, expected_unit))
        if found_unit and found_unit != expected_unit:
            issues.append(
                issue("T13", "ERRO", path, summary["sheet"], "", "", "", material, expected_unit, summary["units"].get(material), "Unidade divergente na aba TOTAL GERAL.")
            )

    steel_sum = sum((sum_material(data_rows, material) for material in materials if material.startswith("O")), ZERO)
    if summary["total_aco"] is not None and not decimal_close(steel_sum, summary["total_aco"]):
        issues.append(
            issue("T13", "ERRO", path, summary["sheet"], "", "", "", "TOTAL ACO", steel_sum, summary["total_aco"], "TOTAL ACO nao bate com soma das bitolas.")
        )

    return issues


def autosize_columns(worksheet: Any, max_width: int = 60) -> None:
    for col_idx in range(1, worksheet.max_column + 1):
        letter = get_column_letter(col_idx)
        width = 10
        for row_idx in range(1, worksheet.max_row + 1):
            value = worksheet.cell(row_idx, col_idx).value
            if value is not None:
                width = max(width, min(max_width, len(str(value)) + 2))
        worksheet.column_dimensions[letter].width = width


def style_header(worksheet: Any, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = f"A{row + 1}"


def write_report(result: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "RESUMO_TESTES"

    summary.append(["TESTE", "NOME", "STATUS", "DETALHE"])
    for test in result["tests"]:
        summary.append([test["teste"], test["nome"], test["status"], test["detalhe"]])
    style_header(summary)
    color_status_cells(summary, 3)
    autosize_columns(summary)

    issues_ws = workbook.create_sheet("DIVERGENCIAS_E_ALERTAS")
    headers = ["teste", "severidade", "arquivo", "aba", "linha", "coluna", "elemento", "material", "esperado", "encontrado", "detalhe"]
    issues_ws.append([header.upper() for header in headers])
    for item in sorted(result["issues"], key=lambda row: (row["teste"], STATUS_ORDER.get(row["severidade"], 9), str(row["linha"]), row["material"])):
        issues_ws.append([item.get(header, "") for header in headers])
    style_header(issues_ws)
    color_status_cells(issues_ws, 2)
    autosize_columns(issues_ws, max_width=80)

    write_normalized_sheet(workbook, "BRUTA_NORMALIZADA", result["raw"], result["materials"])
    write_normalized_sheet(workbook, "TRATADA_NORMALIZADA", result["treated"], result["materials"])
    write_totals_sheet(workbook, result)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def color_status_cells(worksheet: Any, column: int) -> None:
    fills = {
        "OK": PatternFill("solid", fgColor="C6EFCE"),
        "ALERTA": PatternFill("solid", fgColor="FFEB9C"),
        "ERRO": PatternFill("solid", fgColor="FFC7CE"),
    }
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row, column).value
        if value in fills:
            worksheet.cell(row, column).fill = fills[value]


def write_normalized_sheet(workbook: Workbook, title: str, data: dict[str, Any], materials: list[str]) -> None:
    worksheet = workbook.create_sheet(title)
    headers = ["ORIGEM_LINHA", "ELEMENTO", "NIVEL"] + materials
    worksheet.append(headers)
    for row in data["rows"]:
        worksheet.append(
            [row["row"], row["elemento"], row.get("nivel", "")]
            + [decimal_to_report(row["values"].get(material, ZERO)) for material in materials]
        )
    style_header(worksheet)
    autosize_columns(worksheet)


def write_totals_sheet(workbook: Workbook, result: dict[str, Any]) -> None:
    worksheet = workbook.create_sheet("TOTAIS")
    worksheet.append(["MATERIAL", "SOMA_BRUTA_REAL", "SOMA_TRATADA", "TOTAL_GERAL", "TOTAL_TODOS", "UNIDADE"])
    treated_total = result["treated"]["total_rows"][-1] if result["treated"]["total_rows"] else {"values": {}}
    summary = result["treated"]["summary"]
    for material in result["materials"]:
        worksheet.append(
            [
                material,
                decimal_to_report(sum_material(result["raw"]["rows"], material)),
                decimal_to_report(sum_material(result["treated"]["rows"], material)),
                decimal_to_report(summary["totals"].get(material)),
                decimal_to_report(treated_total["values"].get(material, ZERO)),
                summary["units"].get(material, ""),
            ]
        )
    style_header(worksheet)
    autosize_columns(worksheet)


def default_output_path(treated_path: Path) -> Path:
    return treated_path.with_name(treated_path.stem + "_CONFERENCIA.xlsx")


def choose_files_with_tk() -> tuple[Path, Path, Path]:
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.withdraw()
    try:
        root.attributes("-topmost", True)
    except tk.TclError:
        pass

    raw = filedialog.askopenfilename(
        title="1 - Escolha a planilha BRUTA do Navisworks/AltoQi",
        filetypes=[("Excel", "*.xlsx")],
        parent=root,
    )
    if not raw:
        raise SystemExit("Cancelado: planilha bruta.")
    treated = filedialog.askopenfilename(
        title="2 - Escolha a planilha TRATADA pelo codigo",
        filetypes=[("Excel", "*.xlsx")],
        parent=root,
    )
    if not treated:
        raise SystemExit("Cancelado: planilha tratada.")
    output = filedialog.asksaveasfilename(
        title="Salvar relatorio de conferencia",
        initialfile=default_output_path(Path(treated)).name,
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")],
        parent=root,
    )
    if not output:
        raise SystemExit("Cancelado: relatorio.")

    messagebox.showinfo("Conferencia", "Clique OK para iniciar a conferencia.", parent=root)
    root.destroy()
    return Path(raw), Path(treated), Path(output)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Confere planilha bruta do Navisworks contra planilha tratada.")
    parser.add_argument("arquivos", nargs="*", help="Ordem esperada: bruta.xlsx tratada.xlsx")
    parser.add_argument("--saida", "-s", help="Caminho do relatorio .xlsx")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    if len(args.arquivos) == 0:
        raw_path, treated_path, output_path = choose_files_with_tk()
    elif len(args.arquivos) == 2:
        raw_path = Path(args.arquivos[0])
        treated_path = Path(args.arquivos[1])
        output_path = Path(args.saida) if args.saida else default_output_path(treated_path)
    else:
        print("Uso: python conferir_lista_materiais.py bruta.xlsx tratada.xlsx [--saida relatorio.xlsx]")
        return 2

    result = run_tests(raw_path, treated_path)
    write_report(result, output_path)

    status_counts = Counter(test["status"] for test in result["tests"])
    print(f"Relatorio gerado: {output_path}")
    print(f"Testes OK: {status_counts.get('OK', 0)}")
    print(f"Alertas: {status_counts.get('ALERTA', 0)}")
    print(f"Erros: {status_counts.get('ERRO', 0)}")
    if result["swapped"]:
        print("Aviso: os arquivos foram detectados invertidos e corrigidos automaticamente.")
    return 1 if status_counts.get("ERRO", 0) else 0


if __name__ == "__main__":
    raise SystemExit(main())
